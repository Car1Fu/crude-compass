from __future__ import annotations

import csv
import calendar
import hashlib
import re
import sqlite3
import unicodedata
import zipfile
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from datetime import datetime, timedelta
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


ROOT_DIR = Path(__file__).resolve().parent
DATA_DIR = ROOT_DIR / "data"
SAMPLE_DATA_DIR = DATA_DIR / "data_xlsx"
DEFAULT_WORKBOOK_DIR = SAMPLE_DATA_DIR
DEFAULT_DB_PATH = DATA_DIR / "crude_compass.sqlite3"
DEFAULT_SAMPLE_XLSX = SAMPLE_DATA_DIR / "_sample_crude_data.xlsx"
LEGACY_SAMPLE_XLSX = ROOT_DIR / "_sample_crude_data.xlsx"
RAW_MACRO_INDICATOR_CSV = SAMPLE_DATA_DIR / "宏观指标_汇总.csv"
MACRO_SPLIT_SPECS = {
    "月": {
        "file_name": "宏观指标（月度）.csv",
        "dataset_code": "macro_monthly_indicators",
    },
    "日": {
        "file_name": "宏观指标（日度）.csv",
        "dataset_code": "macro_daily_indicators",
    },
}
DEFAULT_SAMPLE_XLSX_CANDIDATES = (
    DEFAULT_SAMPLE_XLSX,
    LEGACY_SAMPLE_XLSX,
)

GENERIC_DATASET_CODES = {
    "原油供需相关数据.xlsx": "country_supply_demand_monthly",
    "OPEC消费量预测值.xlsx": "opec_consumption_forecast",
    "全球供需.xlsx": "global_supply_demand_monthly",
    "opec原油产量.xlsx": "opec_crude_production_monthly",
    "原油需求量预测.xlsx": "crude_demand_forecast_quarterly",
    "\u70bc\u6cb9\u5f00\u5de5\u7387.xlsx": "refinery_utilization_monthly",
    "原油期货价格（日）.xlsx": "crude_futures_daily",
    "原油现货价格（日）.xlsx": "crude_spot_daily",
    "成品油期货价格（日）.xlsx": "refined_futures_daily",
    "成品油柴油现货价格（日）.xlsx": "refined_diesel_spot_daily",
    "成品油汽油现货价格（日）.xlsx": "refined_gasoline_spot_daily",
    "成品油煤油现货价格（日）.xlsx": "refined_kerosene_spot_daily",
    "成品油燃料油现货价格（日）.xlsx": "refined_fuel_oil_spot_daily",
    "成品油石脑油现货价格（日）.xlsx": "refined_naphtha_spot_daily",
    "宏观指标（月度）.csv": "macro_monthly_indicators",
    "宏观指标（日度）.csv": "macro_daily_indicators",
}

GENERIC_HEADER_MAP = {
    "name": "series_name",
    "名称": "series_name",
    "简称": "series_name",
    "date": "trade_date",
    "日期": "trade_date",
    "close": "close_price",
    "收盘价": "close_price",
    "settlement": "settlement_price",
    "结算价": "settlement_price",
    "volume": "volume",
    "成交量": "volume",
    "price": "price",
    "价格": "price",
    "value": "value",
    "值": "value",
    "consumption": "consumption",
    "消费量": "consumption",
    "source": "source_label",
    "来源": "source_label",
    "frequency": "frequency_label",
    "freq": "frequency_label",
    "频率": "frequency_label",
    "unit": "unit_label",
    "单位": "unit_label",
    "series_id": "series_id",
    "indicator_id": "series_id",
    "指标id": "series_id",
}

GENERIC_DIMENSION_KEYS = (
    "series_name",
    "trade_date",
    "source_label",
    "frequency_label",
    "unit_label",
    "series_id",
)
GENERIC_METRIC_KEYS = (
    "close_price",
    "settlement_price",
    "volume",
    "price",
    "value",
    "consumption",
)

GENERIC_PRIMARY_METRIC = {
    "opec_consumption_forecast": "consumption",
    "crude_futures_daily": "close_price",
    "crude_spot_daily": "price",
    "refined_futures_daily": "close_price",
    "refined_diesel_spot_daily": "price",
    "refined_gasoline_spot_daily": "price",
    "refined_kerosene_spot_daily": "price",
    "refined_fuel_oil_spot_daily": "price",
    "refined_naphtha_spot_daily": "price",
    "macro_monthly_indicators": "value",
    "macro_daily_indicators": "value",
    "opec_crude_production_monthly": "value",
    "crude_demand_forecast_quarterly": "value",
    "refinery_utilization_monthly": "value",
    "global_supply_demand_monthly": "value",
    "country_supply_demand_monthly": "production",
}
SPARSE_GENERIC_DATASET_CODES = {
    "macro_monthly_indicators",
    "macro_daily_indicators",
    "country_supply_demand_monthly",
}

SUPPLY_DASHBOARD_DATASET_CODE = "global_supply_demand_monthly"
SUPPLY_DASHBOARD_SERIES = {
    "globalSupply": "world_oil_supply",
    "globalDemand": "world_oil_demand",
    "balance": "supply_demand_gap",
    "floatingStorage": "store",
}
SUPPLY_COUNTRY_DATASET_CODE = "country_supply_demand_monthly"
SUPPLY_COUNTRY_METRIC_COLUMNS = (
    ("production", 3, "Thousand Barrels"),
    ("imports", 4, "Thousand Barrels"),
    ("exports", 5, "Thousand Barrels"),
    ("stock_change", 6, "Thousand Barrels"),
    ("products_supplied", 7, "Thousand Barrels"),
    ("refinery_utilization", 8, "%"),
)
SUPPLY_COUNTRY_CODE_MAP = {
    "U.S.": "US",
    "Canada": "CA",
    "China": "CN",
    "India": "IN",
    "Japan": "JP",
    "Korea": "KR",
    "Saudi Arabia": "SA",
}
SUPPLY_PORT_DATASET_CODE = "ports_reference_all_ports"
SUPPLY_PORT_WORKBOOK_NAMES = {"ports_wpi.xlsx"}
SUPPLY_PORT_DEFAULT_SHEET = "All_Ports"
SUPPLY_PORT_REQUIRED_COLUMNS = (
    "World Port Index Number",
    "Main Port Name",
    "Alternate Port Name",
    "UN/LOCODE",
    "Country Code",
    "World Water Body",
    "Latitude",
    "Longitude",
    "Channel Depth (m)",
    "Anchorage Depth (m)",
    "Cargo Pier Depth (m)",
    "Oil Terminal Depth (m)",
    "Harbor Size",
    "Harbor Type",
    "Shelter Afforded",
)
PORT_MISSING_VALUE_TOKENS = {
    "",
    "-",
    "--",
    "unknown",
    "unk",
    "n/a",
    "na",
    "none",
    "null",
}
SUPPLY_TRACKING_DATASET_CODE = "oceanook_tanker_tracking"
SUPPLY_TRACKING_WORKBOOK_NAMES = {
    "oceanook_tankers_enriched.csv",
    "oceanook_tankers_enriched.xlsx",
}
DEFAULT_SUPPLY_TRACKING_XLSX = SAMPLE_DATA_DIR / "oceanook_tankers_enriched.xlsx"
DEFAULT_SUPPLY_TRACKING_CSV = SAMPLE_DATA_DIR / "oceanook_tankers_enriched.csv"
DEFAULT_SUPPLY_TRACKING_CANDIDATES = (
    DEFAULT_SUPPLY_TRACKING_XLSX,
    DEFAULT_SUPPLY_TRACKING_CSV,
)
NEWS_EXTRACTION_WORKBOOK_NAMES = {"news_extraction.xlsx"}
NEWS_DATASET_CODE = "module_news_feed"
NEWS_THEME_SECTION_MAP = {
    "供给政策（opec+）": ("Supply", "供给政策（OPEC+）"),
    "风险事件（地缘）": ("Geopolitics", "风险事件（地缘）"),
    "数据与流向（库存）": ("Inventory", "数据与流向（库存）"),
    "航运与物价（运价）": ("Freight", "航运与物流（运价）"),
    "结构与炼化（价差）": ("Spreads", "结构与炼化（价差）"),
    "美元与利率（宏观）": ("Macro", "美元与利率（宏观）"),
    "消费与经济（需求）": ("Demand", "消费与经济（需求）"),
}
SUPPLY_TRACKING_HEADER_MAP = {
    "船名": "vessel_name",
    "mmsi": "mmsi",
    "imo": "imo",
    "imonumber": "imo",
    "shipimo": "imo",
    "当时位置": "current_position_raw",
    "起始地": "origin_name",
    "上一港": "previous_port_name",
    "目的地": "destination_name",
    "最大负载量": "dwt",
    "来源船型": "vessel_type",
    "来源_船型": "vessel_type",
    "sourcevesseltype": "vessel_type",
    "shipname": "vessel_name",
    "currentposition": "current_position_raw",
    "origin": "origin_name",
    "previousport": "previous_port_name",
    "destination": "destination_name",
    "maxdwt": "dwt",
}
SUPPLY_TRACKING_REQUIRED_FIELDS = (
    "vessel_name",
    "mmsi",
    "current_position_raw",
    "origin_name",
    "previous_port_name",
    "destination_name",
    "dwt",
    "vessel_type",
)
SUPPLY_TRACKING_NORMALIZED_HEADER_MAP = {
    re.sub(r"[\s_]+", "", key.strip().lower()): value
    for key, value in SUPPLY_TRACKING_HEADER_MAP.items()
}
SUPPLY_TRACKING_PLACEHOLDER_TOKENS = {
    "",
    "-",
    "--",
    "na",
    "n a",
    "n/a",
    "none",
    "null",
    "unknown",
    "pending",
    "china",
    "kuwait china",
    "turkish vessel crew",
}
SUPPLY_TRACKING_SOURCE_LABEL = "Oceanook"
SUPPLY_TRACKING_PORT_ALIAS_MAP = {
    "aberdeen hong kong": "hong kong",
    "botlek netherlands": "rotterdam",
    "brunsbuettel germany": "brunsbuttel canal terminals",
    "changi singapore": "singapore",
    "cochin anch india": "kochi cochin",
    "dardanelles north anch turkey": "canakkale",
    "fos sur mer france": "fos",
    "gibraltar west anch gibraltar": "europa point",
    "gonfreville l orcher france": "port of le havre",
    "kizomba": "kizomba a terminal",
    "lauenburg elbe germany": "brunsbuttel canal terminals",
    "le grand quevilly france": "port of rouen",
    "mina al fahl anch oman": "mina al fahl",
    "mundra term india": "mundra",
    "nowy port gdansk poland": "nowy port",
    "piraeus greece": "piraievs",
    "port said egypt": "port said",
    "ras tanura saudi arabia": "ras tanura",
    "shuaiba anch kuwait": "shuaiba",
    "singapore anch 4 singapore": "singapore",
    "singapore anch 5 singapore": "singapore",
    "tanjung pelepas anch malaysia": "tanjung pelepas",
}
PORT_NAME_ZH_OVERRIDES = {
    "Xiamen": "厦门港",
    "Shanghai": "上海港",
    "Ningbo": "宁波港",
    "Qingdao Gang": "青岛港",
    "Dalian": "大连港",
    "Zhoushan": "舟山港",
    "Tianjin Xingang": "天津新港",
    "Tianjin": "天津港",
    "Guangzhou": "广州港",
    "Shenzhen": "深圳港",
    "Yantai": "烟台港",
    "Rizhao": "日照港",
    "Lianyungang": "连云港",
    "Zhanjiang": "湛江港",
    "Qinzhou": "钦州港",
    "Beihai": "北海港",
    "Fuzhou": "福州港",
    "Quanzhou": "泉州港",
    "Xingang": "新港",
}
PORT_HARBOR_SIZE_ZH = {
    "very small": "极小型",
    "small": "小型",
    "medium": "中型",
    "large": "大型",
    "very large": "特大型",
}
PORT_HARBOR_TYPE_ZH = {
    "river (natural)": "河港（天然）",
    "coastal (breakwater)": "沿海港（防波堤）",
    "coastal (natural)": "沿海港（天然）",
    "open roadstead": "开放式锚地港",
    "river basin": "河湾港",
    "canal": "运河港",
    "lake (natural)": "湖港（天然）",
    "lake (artificial)": "湖港（人工）",
    "coastal (artificial)": "沿海港（人工）",
    "coastal": "沿海港",
    "river": "河港",
}
PORT_SHELTER_AFFORDED_ZH = {
    "excellent": "优良",
    "good": "良好",
    "fair": "一般",
    "poor": "较差",
}
PORT_WATER_BODY_ZH = {
    "North Pacific Ocean": "北太平洋",
    "South Pacific Ocean": "南太平洋",
    "North Atlantic Ocean": "北大西洋",
    "South Atlantic Ocean": "南大西洋",
    "Indian Ocean": "印度洋",
    "Mediterranean Sea": "地中海",
    "Black Sea": "黑海",
    "Red Sea": "红海",
    "Persian Gulf": "波斯湾",
    "Arabian Sea": "阿拉伯海",
    "East China Sea": "东海",
    "South China Sea": "南海",
    "Yellow Sea": "黄海",
    "Bohai Sea": "渤海",
    "Taiwan Strait": "台湾海峡",
    "Sea of Japan": "日本海",
    "Gulf of Mexico": "墨西哥湾",
    "Caribbean Sea": "加勒比海",
}

_XML_NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}

_HEADER_MAP = {
    "variable": "symbol",
    "date": "trade_date",
    "cha": "change",
    "per_cha": "change_pct",
    "open": "open_price",
    "high": "high_price",
    "low": "low_price",
    "close": "close_price",
    "volume": "volume",
}


def ensure_data_dir() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)


def resolve_sample_xlsx(sample_path: Path | str | None = None) -> Path:
    if sample_path is not None:
        return Path(sample_path)

    for candidate in DEFAULT_SAMPLE_XLSX_CANDIDATES:
        if candidate.exists():
            return candidate

    return DEFAULT_SAMPLE_XLSX


def resolve_supply_tracking_workbook(source_path: Path | str | None = None) -> Path:
    if source_path is not None:
        return Path(source_path)

    for candidate in DEFAULT_SUPPLY_TRACKING_CANDIDATES:
        if candidate.exists():
            return candidate

    return DEFAULT_SUPPLY_TRACKING_CSV


def get_connection(db_path: Path | str = DEFAULT_DB_PATH) -> sqlite3.Connection:
    ensure_data_dir()
    connection = sqlite3.connect(Path(db_path))
    connection.row_factory = sqlite3.Row
    return connection


def _ensure_table_columns(
    connection: sqlite3.Connection,
    table_name: str,
    required_columns: dict[str, str],
) -> None:
    existing_columns = {
        row["name"]
        for row in connection.execute(f"PRAGMA table_info({table_name})").fetchall()
    }
    for column_name, column_definition in required_columns.items():
        if column_name in existing_columns:
            continue
        connection.execute(
            f"ALTER TABLE {table_name} ADD COLUMN {column_name} {column_definition}"
        )


def init_market_database(db_path: Path | str = DEFAULT_DB_PATH) -> None:
    with get_connection(db_path) as connection:
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS market_daily_prices (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                symbol TEXT NOT NULL,
                trade_date TEXT NOT NULL,
                change REAL,
                change_pct REAL,
                open_price REAL,
                high_price REAL,
                low_price REAL,
                close_price REAL,
                volume INTEGER,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(symbol, trade_date)
            )
            """
        )
        connection.execute(
            """
            CREATE INDEX IF NOT EXISTS idx_market_daily_prices_symbol_date
            ON market_daily_prices(symbol, trade_date)
            """
        )
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS market_generic_daily_metrics (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                dataset_code TEXT NOT NULL,
                dataset_name TEXT NOT NULL,
                source_file TEXT NOT NULL,
                source_sheet TEXT NOT NULL,
                series_name TEXT NOT NULL,
                source_label TEXT,
                frequency_label TEXT,
                unit_label TEXT,
                series_id TEXT,
                trade_date TEXT NOT NULL,
                metric_key TEXT NOT NULL,
                metric_value REAL NOT NULL,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(dataset_code, series_name, trade_date, metric_key)
            )
            """
        )
        connection.execute(
            """
            CREATE INDEX IF NOT EXISTS idx_market_generic_daily_metrics_lookup
            ON market_generic_daily_metrics(dataset_code, series_name, trade_date, metric_key)
            """
        )
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS port_reference_data (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                dataset_code TEXT NOT NULL,
                dataset_name TEXT NOT NULL,
                source_file TEXT NOT NULL,
                source_sheet TEXT NOT NULL,
                wpi_number TEXT,
                main_port_name TEXT NOT NULL,
                alternate_port_name TEXT,
                unlocode TEXT,
                country_name TEXT,
                world_water_body TEXT,
                latitude REAL NOT NULL,
                longitude REAL NOT NULL,
                channel_depth_m TEXT,
                anchorage_depth_m TEXT,
                cargo_pier_depth_m TEXT,
                oil_terminal_depth_m TEXT,
                harbor_size TEXT,
                harbor_type TEXT,
                shelter_afforded TEXT,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(dataset_code, wpi_number)
            )
            """
        )
        connection.execute(
            """
            CREATE INDEX IF NOT EXISTS idx_port_reference_data_lookup
            ON port_reference_data(dataset_code, main_port_name, country_name)
            """
        )
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS supply_tanker_tracking (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                dataset_code TEXT NOT NULL,
                dataset_name TEXT NOT NULL,
                source_file TEXT NOT NULL,
                source_sheet TEXT NOT NULL,
                source_row_number INTEGER NOT NULL,
                vessel_name TEXT NOT NULL,
                imo TEXT,
                mmsi TEXT,
                vessel_type TEXT NOT NULL,
                dwt REAL,
                current_position_raw TEXT,
                current_latitude REAL NOT NULL,
                current_longitude REAL NOT NULL,
                position_label TEXT,
                origin_name TEXT,
                origin_latitude REAL,
                origin_longitude REAL,
                origin_match_method TEXT,
                origin_matched_port_name TEXT,
                previous_port_name TEXT,
                previous_port_latitude REAL,
                previous_port_longitude REAL,
                previous_port_match_method TEXT,
                previous_port_matched_port_name TEXT,
                destination_name TEXT,
                destination_latitude REAL,
                destination_longitude REAL,
                destination_match_method TEXT,
                destination_matched_port_name TEXT,
                speed REAL,
                heading REAL,
                cargo_status TEXT NOT NULL DEFAULT 'unknown',
                eta TEXT,
                source_label TEXT,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(dataset_code, source_file, source_sheet, source_row_number)
            )
            """
        )
        connection.execute(
            """
            CREATE INDEX IF NOT EXISTS idx_supply_tanker_tracking_lookup
            ON supply_tanker_tracking(dataset_code, vessel_type, vessel_name)
            """
        )
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS market_news_items (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                dataset_code TEXT NOT NULL,
                source_file TEXT NOT NULL,
                source_sheet TEXT NOT NULL,
                article_key TEXT NOT NULL,
                title TEXT NOT NULL,
                source_label TEXT NOT NULL,
                published_at TEXT NOT NULL,
                theme_raw TEXT NOT NULL,
                section_key TEXT NOT NULL,
                section_name TEXT NOT NULL,
                market_impact_score REAL,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(dataset_code, article_key)
            )
            """
        )
        connection.execute(
            """
            CREATE INDEX IF NOT EXISTS idx_market_news_items_section_time
            ON market_news_items(dataset_code, section_key, published_at DESC)
            """
        )
        _ensure_table_columns(
            connection,
            "market_generic_daily_metrics",
            {
                "source_label": "TEXT",
                "frequency_label": "TEXT",
                "unit_label": "TEXT",
                "series_id": "TEXT",
            },
        )


def normalize_symbol(symbol: str | None) -> str | None:
    if symbol is None:
        return None
    normalized = str(symbol).strip().lower()
    if normalized == "wti":
        return "WTI"
    if normalized == "brent":
        return "Brent"
    return None


def _load_shared_strings(zf: zipfile.ZipFile) -> list[str]:
    if "xl/sharedStrings.xml" not in zf.namelist():
        return []

    shared_strings_root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
    shared_strings: list[str] = []
    for item in shared_strings_root.findall("main:si", _XML_NS):
        chunks = [node.text or "" for node in item.findall(".//main:t", _XML_NS)]
        shared_strings.append("".join(chunks))
    return shared_strings


def _sheet_targets(zf: zipfile.ZipFile) -> list[tuple[str, str]]:
    workbook_root = ET.fromstring(zf.read("xl/workbook.xml"))
    rels_root = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    rel_map = {
        rel.attrib["Id"]: rel.attrib["Target"]
        for rel in rels_root.findall("rel:Relationship", _XML_NS)
    }

    targets: list[tuple[str, str]] = []
    for sheet in workbook_root.findall("main:sheets/main:sheet", _XML_NS):
        name = sheet.attrib.get("name", "Sheet1")
        rel_id = sheet.attrib.get(
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
        )
        if not rel_id or rel_id not in rel_map:
            continue
        target = rel_map[rel_id]
        if not target.startswith("xl/"):
            target = f"xl/{target}"
        targets.append((name, target))
    return targets


def _read_cell_value(cell: ET.Element, shared_strings: list[str]) -> str:
    cell_type = cell.attrib.get("t")
    value_node = cell.find("main:v", _XML_NS)
    if value_node is None:
        inline_node = cell.find("main:is", _XML_NS)
        if inline_node is None:
            return ""
        chunks = [node.text or "" for node in inline_node.findall(".//main:t", _XML_NS)]
        return "".join(chunks)

    raw_value = value_node.text or ""
    if cell_type == "s" and raw_value:
        return shared_strings[int(raw_value)]
    return raw_value


def _column_reference_to_index(cell_reference: str) -> int:
    letters = "".join(character for character in cell_reference if character.isalpha())
    index = 0
    for character in letters:
        index = index * 26 + (ord(character.upper()) - 64)
    return max(index - 1, 0)


def _read_sheet_rows(
    zf: zipfile.ZipFile,
    target_path: str,
    shared_strings: list[str],
) -> list[list[str]]:
    sheet_root = ET.fromstring(zf.read(target_path))
    rows: list[list[str]] = []
    for row in sheet_root.findall("main:sheetData/main:row", _XML_NS):
        cells = row.findall("main:c", _XML_NS)
        if not cells:
            rows.append([])
            continue

        max_column_index = max(
            _column_reference_to_index(cell.attrib.get("r", "A1"))
            for cell in cells
        )
        values = [""] * (max_column_index + 1)
        for cell in cells:
            column_index = _column_reference_to_index(cell.attrib.get("r", "A1"))
            values[column_index] = _read_cell_value(cell, shared_strings)

        while values and not str(values[-1]).strip():
            values.pop()
        rows.append(values)

    return rows


def _normalize_header_label(raw_value: str) -> str:
    return str(raw_value or "").strip().lower()


def _excel_date_to_iso(raw_value: str) -> str:
    value = raw_value.strip()
    if not value:
        return ""

    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(value, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue

    try:
        serial = float(value)
    except ValueError:
        return ""
    base = datetime(1899, 12, 30)
    return (base + timedelta(days=serial)).strftime("%Y-%m-%d")


def _to_float(raw_value: str) -> float | None:
    value = raw_value.strip()
    if not value:
        return None
    try:
        return float(value)
    except ValueError:
        return None


def _to_int(raw_value: str) -> int | None:
    value = raw_value.strip()
    if not value:
        return None
    try:
        return int(float(value))
    except ValueError:
        return None


def read_market_rows_from_xlsx(
    excel_path: Path | str,
    sheet_name: str | None = None,
) -> list[dict[str, object]]:
    xlsx_path = Path(excel_path)
    if not xlsx_path.exists():
        raise FileNotFoundError(f"Excel file not found: {xlsx_path}")

    rows: list[dict[str, object]] = []
    with zipfile.ZipFile(xlsx_path) as zf:
        shared_strings = _load_shared_strings(zf)
        sheet_targets = _sheet_targets(zf)
        if sheet_name:
            sheet_targets = [item for item in sheet_targets if item[0] == sheet_name]
            if not sheet_targets:
                raise ValueError(f"Sheet not found: {sheet_name}")

        target_name, target_path = sheet_targets[0]
        raw_rows = _read_sheet_rows(zf, target_path, shared_strings)
        if not raw_rows:
            return rows

        headers = [str(value).strip() for value in raw_rows[0]]
        mapped_headers = [_HEADER_MAP.get(header, header) for header in headers]

        for row in raw_rows[1:]:
            values = [str(value) for value in row]
            if not values:
                continue

            raw_record = dict(zip(mapped_headers, values))
            symbol = normalize_symbol(raw_record.get("symbol"))
            if symbol is None:
                continue

            record = {
                "symbol": symbol,
                "trade_date": _excel_date_to_iso(str(raw_record.get("trade_date", ""))),
                "change": _to_float(str(raw_record.get("change", ""))),
                "change_pct": _to_float(str(raw_record.get("change_pct", ""))),
                "open_price": _to_float(str(raw_record.get("open_price", ""))),
                "high_price": _to_float(str(raw_record.get("high_price", ""))),
                "low_price": _to_float(str(raw_record.get("low_price", ""))),
                "close_price": _to_float(str(raw_record.get("close_price", ""))),
                "volume": _to_int(str(raw_record.get("volume", ""))),
                "source_sheet": target_name,
            }
            if record["trade_date"]:
                rows.append(record)

    return rows


def _resolve_dataset_code(workbook_path: Path) -> str:
    if _is_port_reference_workbook(workbook_path):
        return SUPPLY_PORT_DATASET_CODE

    mapped = GENERIC_DATASET_CODES.get(workbook_path.name)
    if mapped:
        return mapped

    ascii_slug = re.sub(r"[^0-9a-zA-Z]+", "_", workbook_path.stem).strip("_").lower()
    if ascii_slug:
        return ascii_slug

    digest = hashlib.sha1(workbook_path.name.encode("utf-8")).hexdigest()[:10]
    return f"dataset_{digest}"


def list_additional_market_workbooks(
    workbook_dir: Path | str = DEFAULT_WORKBOOK_DIR,
) -> list[Path]:
    directory = Path(workbook_dir)
    if not directory.exists():
        return []

    workbook_paths = [
        path
        for path in directory.glob("*.xlsx")
        if path.name != DEFAULT_SAMPLE_XLSX.name and not path.name.startswith("~$")
    ]
    workbook_paths.extend(
        path
        for path in directory.glob("*.csv")
        if path.name.lower() in SUPPLY_TRACKING_WORKBOOK_NAMES
    )
    return sorted(workbook_paths)


def _read_csv_rows(csv_path: Path) -> tuple[list[list[str]], str]:
    last_error: UnicodeDecodeError | None = None
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
        try:
            with csv_path.open("r", encoding=encoding, newline="") as handle:
                return list(csv.reader(handle)), encoding
        except UnicodeDecodeError as error:
            last_error = error
            continue
    if last_error is not None:
        raise last_error
    return [], "utf-8-sig"


def _normalize_source_label(raw_value: str) -> str:
    normalized = str(raw_value or "").strip()
    if not normalized:
        return "Wind"
    replacements = {
        "美国劳工部": "美国劳工统计局（BLS）",
        "美联储": "美联储",
        "国家统计局": "国家统计局",
        "Wind": "Wind",
        "行情": "Wind",
        "根据新闻整理": "Wind",
    }
    return replacements.get(normalized, normalized)


def _is_port_reference_workbook(workbook_path: Path) -> bool:
    return workbook_path.name.lower() in SUPPLY_PORT_WORKBOOK_NAMES


def _is_news_extraction_workbook(workbook_path: Path) -> bool:
    return workbook_path.name.lower() in NEWS_EXTRACTION_WORKBOOK_NAMES


def _normalize_news_theme(raw_theme: str | None) -> tuple[str, str, str]:
    theme_text = str(raw_theme or "").strip()
    normalized_theme = theme_text.lower()
    section_key, section_name = NEWS_THEME_SECTION_MAP.get(
        normalized_theme,
        ("Macro", theme_text or "未分类新闻"),
    )
    return theme_text, section_key, section_name


def _normalize_news_publish_time(raw_value: object) -> str:
    if isinstance(raw_value, datetime):
        return raw_value.strftime("%Y-%m-%d %H:%M")

    if raw_value is None:
        return ""

    text = str(raw_value).strip()
    if not text:
        return ""

    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y/%m/%d %H:%M:%S",
        "%Y/%m/%d %H:%M",
        "%Y-%m-%d",
        "%Y/%m/%d",
    ):
        try:
            parsed = datetime.strptime(text, fmt)
            return parsed.strftime("%Y-%m-%d %H:%M")
        except ValueError:
            continue
    return text


def read_news_extraction_workbook(
    excel_path: Path | str,
    sheet_name: str | None = None,
) -> dict[str, object]:
    workbook_path = Path(excel_path)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    target_sheet_name = sheet_name or workbook.sheetnames[0]
    worksheet = workbook[target_sheet_name]

    rows_iter = worksheet.iter_rows(values_only=True)
    try:
        raw_headers = next(rows_iter)
    except StopIteration:
        raw_headers = ()
    headers = [str(header or "").strip() for header in raw_headers]

    required_headers = {
        "chinese_title",
        "source",
        "publish_time",
        "theme",
        "market_impact_score",
    }
    missing_headers = sorted(required_headers.difference(headers))
    if missing_headers:
        missing_text = ", ".join(missing_headers)
        raise ValueError(
            f"News workbook is missing required columns: {missing_text}"
        )

    rows: list[dict[str, object]] = []
    dropped_row_count = 0
    for row_index, values in enumerate(rows_iter, start=2):
        record = dict(zip(headers, values))
        title = str(record.get("chinese_title") or "").strip()
        source_label = str(record.get("source") or "").strip() or "未知来源"
        published_at = _normalize_news_publish_time(record.get("publish_time"))
        theme_raw, section_key, section_name = _normalize_news_theme(record.get("theme"))
        impact_score = _to_float(str(record.get("market_impact_score") or ""))
        if not title or not published_at or not theme_raw:
            dropped_row_count += 1
            continue

        article_key = hashlib.sha1(
            f"{title}|{source_label}|{published_at}|{theme_raw}".encode("utf-8")
        ).hexdigest()[:16]
        rows.append(
            {
                "dataset_code": NEWS_DATASET_CODE,
                "source_file": workbook_path.name,
                "source_sheet": target_sheet_name,
                "article_key": article_key,
                "title": title,
                "source_label": source_label,
                "published_at": published_at,
                "theme_raw": theme_raw,
                "section_key": section_key,
                "section_name": section_name,
                "market_impact_score": impact_score,
                "source_row_number": row_index,
            }
        )

    return {
        "excel_path": str(workbook_path),
        "dataset_code": NEWS_DATASET_CODE,
        "source_sheet": target_sheet_name,
        "source_row_count": max(worksheet.max_row - 1, 0),
        "clean_row_count": len(rows),
        "dropped_row_count": dropped_row_count,
        "rows": rows,
    }


def replace_news_rows(
    dataset_code: str,
    rows: Iterable[dict[str, object]],
    db_path: Path | str = DEFAULT_DB_PATH,
) -> int:
    materialized_rows = list(rows)
    init_market_database(db_path)
    with get_connection(db_path) as connection:
        connection.execute(
            """
            DELETE FROM market_news_items
            WHERE dataset_code = ?
            """,
            (dataset_code,),
        )
        if not materialized_rows:
            return 0
        connection.executemany(
            """
            INSERT INTO market_news_items (
                dataset_code,
                source_file,
                source_sheet,
                article_key,
                title,
                source_label,
                published_at,
                theme_raw,
                section_key,
                section_name,
                market_impact_score
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(dataset_code, article_key) DO UPDATE SET
                source_file = excluded.source_file,
                source_sheet = excluded.source_sheet,
                title = excluded.title,
                source_label = excluded.source_label,
                published_at = excluded.published_at,
                theme_raw = excluded.theme_raw,
                section_key = excluded.section_key,
                section_name = excluded.section_name,
                market_impact_score = excluded.market_impact_score
            """,
            [
                (
                    row["dataset_code"],
                    row["source_file"],
                    row["source_sheet"],
                    row["article_key"],
                    row["title"],
                    row["source_label"],
                    row["published_at"],
                    row["theme_raw"],
                    row["section_key"],
                    row["section_name"],
                    row.get("market_impact_score"),
                )
                for row in materialized_rows
            ],
        )
    return len(materialized_rows)


def import_news_extraction_workbook(
    excel_path: Path | str,
    db_path: Path | str = DEFAULT_DB_PATH,
    sheet_name: str | None = None,
) -> dict[str, object]:
    workbook = read_news_extraction_workbook(excel_path, sheet_name=sheet_name)
    inserted_rows = replace_news_rows(
        workbook["dataset_code"],
        workbook["rows"],
        db_path=db_path,
    )
    return {
        "excel_path": workbook["excel_path"],
        "db_path": str(Path(db_path)),
        "dataset_code": workbook["dataset_code"],
        "source_sheet": workbook["source_sheet"],
        "source_row_count": workbook["source_row_count"],
        "clean_row_count": workbook["clean_row_count"],
        "dropped_row_count": workbook["dropped_row_count"],
        "news_row_count": inserted_rows,
    }


def _is_port_placeholder_value(raw_value: str | None) -> bool:
    value = str(raw_value or "").strip()
    if not value:
        return True
    if value.lower() in PORT_MISSING_VALUE_TOKENS:
        return True
    numeric_value = _to_float(value)
    return numeric_value is not None and abs(numeric_value) < 1e-9


def _format_numeric_display(raw_value: str | None) -> str:
    numeric_value = _to_float(str(raw_value or ""))
    if numeric_value is None:
        return str(raw_value or "").strip()
    return f"{numeric_value:.1f}".rstrip("0").rstrip(".")


def _format_port_depth_display(raw_value: str | None) -> str:
    if _is_port_placeholder_value(raw_value):
        return "待核实"
    return f"{_format_numeric_display(raw_value)} 米"


def _translate_port_name(
    main_port_name: str | None,
    alternate_port_name: str | None = None,
) -> str:
    candidates = [str(main_port_name or "").strip(), str(alternate_port_name or "").strip()]
    for candidate in candidates:
        if not candidate:
            continue
        translated = PORT_NAME_ZH_OVERRIDES.get(candidate)
        if translated:
            return translated
    return candidates[0] if candidates and candidates[0] else "待核实"


def _translate_port_size(raw_value: str | None) -> str:
    value = str(raw_value or "").strip()
    if _is_port_placeholder_value(value):
        return "待核实"
    return PORT_HARBOR_SIZE_ZH.get(value.lower(), value)


def _translate_port_type(raw_value: str | None) -> str:
    value = str(raw_value or "").strip()
    if _is_port_placeholder_value(value):
        return "待核实"

    exact_match = PORT_HARBOR_TYPE_ZH.get(value.lower())
    if exact_match:
        return exact_match

    translated = value
    replacements = (
        ("Open Roadstead", "开放式锚地港"),
        ("Coastal", "沿海港"),
        ("River", "河港"),
        ("Lake", "湖港"),
        ("Canal", "运河港"),
        ("Breakwater", "防波堤"),
        ("Natural", "天然"),
        ("Artificial", "人工"),
        ("Roadstead", "锚地"),
        ("Basin", "港池"),
        ("Marina", "游艇港"),
    )
    for english, chinese in replacements:
        translated = translated.replace(english, chinese)
    translated = translated.replace("(", "（").replace(")", "）")
    return translated


def _translate_port_shelter(raw_value: str | None) -> str:
    value = str(raw_value or "").strip()
    if _is_port_placeholder_value(value):
        return "待核实"
    return PORT_SHELTER_AFFORDED_ZH.get(value.lower(), value)


def _translate_port_water_body(raw_value: str | None) -> str:
    value = str(raw_value or "").strip()
    if _is_port_placeholder_value(value):
        return "待核实"

    parts = [part.strip() for part in value.split(";") if part.strip()]
    translated_parts = [PORT_WATER_BODY_ZH.get(part, part) for part in parts]
    return "；".join(translated_parts) if translated_parts else value


def _port_row_value(port_row: sqlite3.Row | dict[str, object], key: str) -> object | None:
    if isinstance(port_row, sqlite3.Row):
        return port_row[key] if key in port_row.keys() else None
    return port_row.get(key)


def _build_port_detail_label(port_row: sqlite3.Row | dict[str, object]) -> str | None:
    harbor_type = _translate_port_type(_port_row_value(port_row, "harbor_type"))
    harbor_size = _translate_port_size(_port_row_value(port_row, "harbor_size"))
    detail_parts = [
        value
        for value in (harbor_type, harbor_size)
        if value and value != "待核实"
    ]
    if not detail_parts:
        return None
    return " · ".join(detail_parts[:2])


def split_macro_indicator_csv(
    source_path: Path | str = RAW_MACRO_INDICATOR_CSV,
    output_dir: Path | str = DEFAULT_WORKBOOK_DIR,
) -> list[dict[str, object]]:
    source_file = Path(source_path)
    if not source_file.exists():
        return []

    rows, _encoding = _read_csv_rows(source_file)
    if len(rows) < 6:
        raise ValueError(f"Macro indicator csv is malformed: {source_file}")

    output_directory = Path(output_dir)
    output_directory.mkdir(parents=True, exist_ok=True)

    name_row = rows[0]
    frequency_row = rows[1]
    unit_row = rows[2]
    series_id_row = rows[3]
    source_row = rows[4]
    data_rows = rows[5:]

    summaries: list[dict[str, object]] = []
    for frequency_key, spec in MACRO_SPLIT_SPECS.items():
        column_indexes = [
            index
            for index in range(1, len(name_row))
            if index < len(frequency_row)
            and str(frequency_row[index]).strip() == frequency_key
        ]
        output_path = output_directory / str(spec["file_name"])
        written_rows = 0
        with output_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(["名称", "日期", "值", "来源", "频率", "单位", "指标ID"])
            for column_index in column_indexes:
                series_name = str(name_row[column_index]).strip()
                source_label = (
                    _normalize_source_label(str(source_row[column_index]).strip())
                    if column_index < len(source_row)
                    else "Wind"
                )
                frequency_label = (
                    str(frequency_row[column_index]).strip()
                    if column_index < len(frequency_row)
                    else ""
                )
                unit_label = (
                    str(unit_row[column_index]).strip()
                    if column_index < len(unit_row)
                    else ""
                )
                series_id = (
                    str(series_id_row[column_index]).strip()
                    if column_index < len(series_id_row)
                    else ""
                )
                for data_row in data_rows:
                    trade_date = _excel_date_to_iso(str(data_row[0]).strip() if data_row else "")
                    raw_value = (
                        str(data_row[column_index]).strip()
                        if column_index < len(data_row)
                        else ""
                    )
                    if not series_name or not trade_date or not raw_value:
                        continue
                    writer.writerow(
                        [
                            series_name,
                            trade_date,
                            raw_value,
                            source_label,
                            frequency_label,
                            unit_label,
                            series_id,
                        ]
                    )
                    written_rows += 1

        summaries.append(
            {
                "source_csv": str(source_file),
                "output_path": str(output_path),
                "dataset_code": spec["dataset_code"],
                "frequency": frequency_key,
                "series_count": len(column_indexes),
                "row_count": written_rows,
            }
        )

    return summaries


def _deduplicate_port_rows(
    rows: list[dict[str, object]],
) -> tuple[list[dict[str, object]], int]:
    deduplicated_rows: list[dict[str, object]] = []
    seen_keys: set[tuple[object, ...]] = set()
    duplicate_count = 0

    for row in rows:
        dedupe_key = (
            row.get("dataset_code"),
            row.get("wpi_number") or "",
            row.get("main_port_name") or "",
            row.get("latitude"),
            row.get("longitude"),
        )
        if dedupe_key in seen_keys:
            duplicate_count += 1
            continue
        seen_keys.add(dedupe_key)
        deduplicated_rows.append(row)

    return deduplicated_rows, duplicate_count


def read_port_reference_workbook(
    excel_path: Path | str,
    sheet_name: str | None = None,
) -> dict[str, object]:
    workbook_path = Path(excel_path)
    if not workbook_path.exists():
        raise FileNotFoundError(f"Excel file not found: {workbook_path}")
    if not _is_port_reference_workbook(workbook_path):
        raise ValueError(f"Workbook is not mapped as a port reference source: {workbook_path.name}")

    with zipfile.ZipFile(workbook_path) as zf:
        shared_strings = _load_shared_strings(zf)
        sheet_targets = _sheet_targets(zf)
        selected_sheet_name = sheet_name or SUPPLY_PORT_DEFAULT_SHEET
        sheet_targets = [item for item in sheet_targets if item[0] == selected_sheet_name]
        if not sheet_targets:
            raise ValueError(f"Sheet not found: {selected_sheet_name}")

        target_name, target_path = sheet_targets[0]
        sheet_rows = _read_sheet_rows(zf, target_path, shared_strings)
        if len(sheet_rows) < 4:
            raise ValueError(f"Port reference sheet is malformed: {target_name}")

        header_row = [str(value).strip() for value in sheet_rows[2]]
        header_indexes = {
            header_name: index
            for index, header_name in enumerate(header_row)
            if header_name
        }
        missing_headers = [
            column_name
            for column_name in SUPPLY_PORT_REQUIRED_COLUMNS
            if column_name not in header_indexes
        ]
        if missing_headers:
            raise ValueError(
                "Port reference sheet is missing required columns: "
                + ", ".join(missing_headers)
            )

        parsed_rows: list[dict[str, object]] = []
        for row in sheet_rows[3:]:
            if not any(str(value).strip() for value in row):
                continue

            def read_value(column_name: str) -> str:
                index = header_indexes[column_name]
                return str(row[index]).strip() if index < len(row) else ""

            main_port_name = read_value("Main Port Name")
            latitude = _to_float(read_value("Latitude"))
            longitude = _to_float(read_value("Longitude"))
            if not main_port_name or latitude is None or longitude is None:
                continue

            parsed_rows.append(
                {
                    "dataset_code": SUPPLY_PORT_DATASET_CODE,
                    "dataset_name": workbook_path.stem,
                    "source_file": workbook_path.name,
                    "source_sheet": target_name,
                    "wpi_number": read_value("World Port Index Number"),
                    "main_port_name": main_port_name,
                    "alternate_port_name": read_value("Alternate Port Name"),
                    "unlocode": read_value("UN/LOCODE"),
                    "country_name": read_value("Country Code"),
                    "world_water_body": read_value("World Water Body"),
                    "latitude": float(latitude),
                    "longitude": float(longitude),
                    "channel_depth_m": read_value("Channel Depth (m)"),
                    "anchorage_depth_m": read_value("Anchorage Depth (m)"),
                    "cargo_pier_depth_m": read_value("Cargo Pier Depth (m)"),
                    "oil_terminal_depth_m": read_value("Oil Terminal Depth (m)"),
                    "harbor_size": read_value("Harbor Size"),
                    "harbor_type": read_value("Harbor Type"),
                    "shelter_afforded": read_value("Shelter Afforded"),
                }
            )

    deduplicated_rows, duplicate_row_count = _deduplicate_port_rows(parsed_rows)
    return {
        "dataset_code": SUPPLY_PORT_DATASET_CODE,
        "dataset_name": workbook_path.stem,
        "excel_path": str(workbook_path),
        "source_sheet": target_name,
        "source_row_count": len(parsed_rows),
        "duplicate_row_count": duplicate_row_count,
        "clean_row_count": len(deduplicated_rows),
        "series_count": len(deduplicated_rows),
        "rows": deduplicated_rows,
    }


def replace_port_reference_rows(
    dataset_code: str,
    rows: Iterable[dict[str, object]],
    db_path: Path | str = DEFAULT_DB_PATH,
) -> int:
    materialized_rows = list(rows)
    init_market_database(db_path)
    with get_connection(db_path) as connection:
        connection.execute(
            """
            DELETE FROM port_reference_data
            WHERE dataset_code = ?
            """,
            (dataset_code,),
        )
        payload = [
            (
                row["dataset_code"],
                row["dataset_name"],
                row["source_file"],
                row["source_sheet"],
                row.get("wpi_number"),
                row["main_port_name"],
                row.get("alternate_port_name"),
                row.get("unlocode"),
                row.get("country_name"),
                row.get("world_water_body"),
                row["latitude"],
                row["longitude"],
                row.get("channel_depth_m"),
                row.get("anchorage_depth_m"),
                row.get("cargo_pier_depth_m"),
                row.get("oil_terminal_depth_m"),
                row.get("harbor_size"),
                row.get("harbor_type"),
                row.get("shelter_afforded"),
            )
            for row in materialized_rows
        ]

        if payload:
            connection.executemany(
                """
                INSERT INTO port_reference_data (
                    dataset_code,
                    dataset_name,
                    source_file,
                    source_sheet,
                    wpi_number,
                    main_port_name,
                    alternate_port_name,
                    unlocode,
                    country_name,
                    world_water_body,
                    latitude,
                    longitude,
                    channel_depth_m,
                    anchorage_depth_m,
                    cargo_pier_depth_m,
                    oil_terminal_depth_m,
                    harbor_size,
                    harbor_type,
                    shelter_afforded
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(dataset_code, wpi_number) DO UPDATE SET
                    dataset_name = excluded.dataset_name,
                    source_file = excluded.source_file,
                    source_sheet = excluded.source_sheet,
                    main_port_name = excluded.main_port_name,
                    alternate_port_name = excluded.alternate_port_name,
                    unlocode = excluded.unlocode,
                    country_name = excluded.country_name,
                    world_water_body = excluded.world_water_body,
                    latitude = excluded.latitude,
                    longitude = excluded.longitude,
                    channel_depth_m = excluded.channel_depth_m,
                    anchorage_depth_m = excluded.anchorage_depth_m,
                    cargo_pier_depth_m = excluded.cargo_pier_depth_m,
                    oil_terminal_depth_m = excluded.oil_terminal_depth_m,
                    harbor_size = excluded.harbor_size,
                    harbor_type = excluded.harbor_type,
                    shelter_afforded = excluded.shelter_afforded
                """,
                payload,
            )
        inserted_row = connection.execute(
            """
            SELECT COUNT(*) AS total
            FROM port_reference_data
            WHERE dataset_code = ?
            """,
            (dataset_code,),
        ).fetchone()

    return int(inserted_row["total"]) if inserted_row else 0


def import_port_reference_workbook(
    excel_path: Path | str,
    db_path: Path | str = DEFAULT_DB_PATH,
    sheet_name: str | None = None,
) -> dict[str, object]:
    workbook = read_port_reference_workbook(excel_path, sheet_name=sheet_name)
    inserted_port_rows = replace_port_reference_rows(
        workbook["dataset_code"],
        workbook["rows"],
        db_path=db_path,
    )

    return {
        "excel_path": workbook["excel_path"],
        "db_path": str(Path(db_path)),
        "dataset_code": workbook["dataset_code"],
        "dataset_name": workbook["dataset_name"],
        "source_sheet": workbook["source_sheet"],
        "series_count": workbook["series_count"],
        "source_row_count": workbook["source_row_count"],
        "deduplicated_row_count": workbook["clean_row_count"],
        "duplicate_row_count": workbook["duplicate_row_count"],
        "clean_row_count": workbook["clean_row_count"],
        "metric_row_count": inserted_port_rows,
        "dropped_date_count": 0,
        "dropped_dates": [],
        "metric_keys": ["port_reference"],
        "port_row_count": inserted_port_rows,
    }


def _find_generic_sheet(
    zf: zipfile.ZipFile,
    shared_strings: list[str],
    sheet_name: str | None = None,
) -> tuple[str, list[list[str]], int, list[str]]:
    sheet_targets = _sheet_targets(zf)
    if sheet_name:
        sheet_targets = [item for item in sheet_targets if item[0] == sheet_name]
        if not sheet_targets:
            raise ValueError(f"Sheet not found: {sheet_name}")

    for target_name, target_path in sheet_targets:
        sheet_rows = _read_sheet_rows(zf, target_path, shared_strings)
        for header_index, row in enumerate(sheet_rows[:10]):
            mapped_headers = [
                GENERIC_HEADER_MAP.get(_normalize_header_label(value), "")
                for value in row
            ]
            if "series_name" in mapped_headers and "trade_date" in mapped_headers:
                metric_keys = [
                    header
                    for header in mapped_headers
                    if header and header not in GENERIC_DIMENSION_KEYS
                ]
                if metric_keys:
                    return target_name, sheet_rows, header_index, mapped_headers

    if sheet_name:
        raise ValueError(f"No structured data header found in sheet: {sheet_name}")
    raise ValueError("No structured data sheet found in workbook")


def _parse_wide_metadata_workbook(
    workbook_path: Path,
    zf: zipfile.ZipFile,
    shared_strings: list[str],
    sheet_name: str | None = None,
) -> dict[str, object]:
    sheet_targets = _sheet_targets(zf)
    if sheet_name:
        sheet_targets = [item for item in sheet_targets if item[0] == sheet_name]
        if not sheet_targets:
            raise ValueError(f"Sheet not found: {sheet_name}")

    meta_aliases = {
        "指标名称": "series_name",
        "名称": "series_name",
        "频率": "frequency_label",
        "单位": "unit_label",
        "来源": "source_label",
        "指标id": "series_id",
    }

    for target_name, target_path in sheet_targets:
        sheet_rows = _read_sheet_rows(zf, target_path, shared_strings)
        if not sheet_rows:
            continue

        meta_rows: dict[str, list[str]] = {}
        data_start_index = 0
        for index, row in enumerate(sheet_rows[:8]):
            if not row:
                continue
            first_value = str(row[0]).strip()
            alias = meta_aliases.get(first_value) or meta_aliases.get(_normalize_header_label(first_value))
            if alias:
                meta_rows[alias] = row
                data_start_index = max(data_start_index, index + 1)

        name_row = meta_rows.get("series_name")
        if not name_row or len(name_row) <= 1:
            continue

        parsed_rows: list[dict[str, object]] = []
        metric_keys = ["value"]
        source_row = meta_rows.get("source_label", [])
        frequency_row = meta_rows.get("frequency_label", [])
        unit_row = meta_rows.get("unit_label", [])
        series_id_row = meta_rows.get("series_id", [])

        for row in sheet_rows[data_start_index:]:
            if not any(str(value).strip() for value in row):
                continue
            trade_date = _excel_date_to_iso(str(row[0]).strip() if row else "")
            if not trade_date:
                continue

            for column_index in range(1, len(name_row)):
                series_name = str(name_row[column_index]).strip() if column_index < len(name_row) else ""
                if not series_name:
                    continue
                raw_value = str(row[column_index]).strip() if column_index < len(row) else ""
                if not raw_value:
                    continue
                record = _build_generic_record(
                    workbook_path,
                    source_sheet=target_name,
                    metric_keys=metric_keys,
                    raw_record={
                        "series_name": series_name,
                        "trade_date": trade_date,
                        "value": raw_value,
                        "source_label": _normalize_source_label(
                            str(source_row[column_index]).strip()
                            if column_index < len(source_row)
                            else ""
                        ),
                        "frequency_label": (
                            str(frequency_row[column_index]).strip()
                            if column_index < len(frequency_row)
                            else ""
                        ),
                        "unit_label": (
                            str(unit_row[column_index]).strip()
                            if column_index < len(unit_row)
                            else ""
                        ),
                        "series_id": (
                            str(series_id_row[column_index]).strip()
                            if column_index < len(series_id_row)
                            else ""
                        ),
                    },
                )
                if record is not None:
                    parsed_rows.append(record)

        if parsed_rows:
            parsed_rows.sort(
                key=lambda row: (
                    str(row.get("series_name", "")),
                    str(row.get("trade_date", "")),
                )
            )
            return _finalize_generic_market_payload(
                workbook_path,
                target_name,
                metric_keys,
                parsed_rows,
            )

    if sheet_name:
        raise ValueError(f"No wide metadata sheet found in sheet: {sheet_name}")
    raise ValueError("No wide metadata sheet found in workbook")


def _infer_series_frequency_label(trade_dates: list[str]) -> str:
    parsed_dates = []
    for trade_date in trade_dates[:4]:
        try:
            parsed_dates.append(datetime.strptime(trade_date, "%Y-%m-%d"))
        except ValueError:
            continue
    if len(parsed_dates) < 2:
        return ""

    day_deltas = [
        abs((right - left).days)
        for left, right in zip(parsed_dates, parsed_dates[1:])
    ]
    if not day_deltas:
        return ""

    average_delta = sum(day_deltas) / len(day_deltas)
    if 27 <= average_delta <= 35:
        return "月"
    if 80 <= average_delta <= 100:
        return "季"
    if 360 <= average_delta <= 370:
        return "年"
    if 6 <= average_delta <= 8:
        return "周"
    return "日"


def _default_compact_source_label(workbook_path: Path) -> str:
    if _resolve_dataset_code(workbook_path) == SUPPLY_DASHBOARD_DATASET_CODE:
        return "IEA"
    return "Wind"


def _is_missing_metric_value(raw_value: str) -> bool:
    return str(raw_value or "").strip() in {"", "-", "—", "N/A", "n/a", "None", "null"}


def _parse_supply_country_workbook(
    workbook_path: Path,
    zf: zipfile.ZipFile,
    shared_strings: list[str],
    sheet_name: str | None = None,
) -> dict[str, object]:
    if _resolve_dataset_code(workbook_path) != SUPPLY_COUNTRY_DATASET_CODE:
        raise ValueError("Workbook is not mapped to the supply country dataset")

    sheet_targets = _sheet_targets(zf)
    if sheet_name:
        sheet_targets = [item for item in sheet_targets if item[0] == sheet_name]
        if not sheet_targets:
            raise ValueError(f"Sheet not found: {sheet_name}")

    for target_name, target_path in sheet_targets:
        sheet_rows = _read_sheet_rows(zf, target_path, shared_strings)
        if len(sheet_rows) < 3:
            continue

        header_row = [str(value).strip() for value in sheet_rows[0]]
        if len(header_row) < 8:
            continue
        if _normalize_header_label(header_row[1]) != "source":
            continue
        if _normalize_header_label(header_row[2]) != "month":
            continue
        if _normalize_header_label(header_row[3]) != "crude oil production":
            continue
        if _normalize_header_label(header_row[7]) != "petroleum products supplied":
            continue

        unit_row = [str(value).strip() for value in sheet_rows[1]]
        metric_keys = [metric_key for metric_key, _column_index, _unit in SUPPLY_COUNTRY_METRIC_COLUMNS]
        parsed_rows: list[dict[str, object]] = []

        for row in sheet_rows[2:]:
            if not any(str(value).strip() for value in row):
                continue

            series_name = str(row[0]).strip() if row else ""
            source_label = _normalize_source_label(str(row[1]).strip() if len(row) > 1 else "")
            trade_date = _excel_date_to_iso(str(row[2]).strip() if len(row) > 2 else "")
            if not series_name or not trade_date:
                continue

            for metric_key, column_index, fallback_unit in SUPPLY_COUNTRY_METRIC_COLUMNS:
                raw_value = str(row[column_index]).strip() if column_index < len(row) else ""
                metric_value = None if _is_missing_metric_value(raw_value) else _to_float(raw_value)
                parsed_rows.append(
                    {
                        "dataset_code": SUPPLY_COUNTRY_DATASET_CODE,
                        "dataset_name": workbook_path.stem,
                        "source_file": workbook_path.name,
                        "source_sheet": target_name,
                        "series_name": series_name,
                        "source_label": source_label or None,
                        "frequency_label": "月",
                        "unit_label": (
                            unit_row[column_index]
                            if column_index < len(unit_row) and unit_row[column_index]
                            else fallback_unit
                        ),
                        "series_id": f"{series_name}:{metric_key}",
                        "trade_date": trade_date,
                        "metrics": {metric_key: metric_value},
                    }
                )

        if parsed_rows:
            return _finalize_generic_market_payload(
                workbook_path,
                target_name,
                metric_keys,
                parsed_rows,
            )

    if sheet_name:
        raise ValueError(f"No supply country data sheet found in sheet: {sheet_name}")
    raise ValueError("No supply country data sheet found in workbook")


def _parse_compact_wide_workbook(
    workbook_path: Path,
    zf: zipfile.ZipFile,
    shared_strings: list[str],
    sheet_name: str | None = None,
) -> dict[str, object]:
    sheet_targets = _sheet_targets(zf)
    if sheet_name:
        sheet_targets = [item for item in sheet_targets if item[0] == sheet_name]
        if not sheet_targets:
            raise ValueError(f"Sheet not found: {sheet_name}")

    for target_name, target_path in sheet_targets:
        sheet_rows = _read_sheet_rows(zf, target_path, shared_strings)
        if len(sheet_rows) < 3:
            continue

        name_row = [str(value).strip() for value in sheet_rows[0]]
        if len(name_row) <= 1 or _excel_date_to_iso(name_row[0]):
            continue

        first_data_index = None
        for row_index, row in enumerate(sheet_rows[1:8], start=1):
            if row and _excel_date_to_iso(str(row[0]).strip()):
                first_data_index = row_index
                break
        if first_data_index is None:
            continue

        unit_row = sheet_rows[1] if first_data_index > 1 else []
        source_label = _default_compact_source_label(workbook_path)
        metric_keys = ["value"]
        parsed_rows: list[dict[str, object]] = []
        detected_trade_dates: list[str] = []

        for row in sheet_rows[first_data_index:]:
            if not any(str(value).strip() for value in row):
                continue

            trade_date = _excel_date_to_iso(str(row[0]).strip() if row else "")
            if not trade_date:
                continue
            detected_trade_dates.append(trade_date)

            for column_index in range(1, len(name_row)):
                series_name = str(name_row[column_index]).strip() if column_index < len(name_row) else ""
                if not series_name:
                    continue
                raw_value = str(row[column_index]).strip() if column_index < len(row) else ""
                if not raw_value:
                    continue

                record = _build_generic_record(
                    workbook_path,
                    source_sheet=target_name,
                    metric_keys=metric_keys,
                    raw_record={
                        "series_name": series_name,
                        "trade_date": trade_date,
                        "value": raw_value,
                        "source_label": source_label,
                        "frequency_label": "",
                        "unit_label": (
                            str(unit_row[column_index]).strip()
                            if column_index < len(unit_row)
                            else ""
                        ),
                        "series_id": "",
                    },
                )
                if record is not None:
                    parsed_rows.append(record)

        if parsed_rows:
            frequency_label = _infer_series_frequency_label(detected_trade_dates)
            for row in parsed_rows:
                if not row.get("frequency_label"):
                    row["frequency_label"] = frequency_label or None
            parsed_rows.sort(
                key=lambda row: (
                    str(row.get("series_name", "")),
                    str(row.get("trade_date", "")),
                )
            )
            return _finalize_generic_market_payload(
                workbook_path,
                target_name,
                metric_keys,
                parsed_rows,
            )

    if sheet_name:
        raise ValueError(f"No compact wide sheet found in sheet: {sheet_name}")
    raise ValueError("No compact wide sheet found in workbook")


def _drop_incomplete_dates(
    rows: list[dict[str, object]],
    *,
    metric_keys: Iterable[str],
) -> tuple[list[dict[str, object]], list[str]]:
    metric_key_list = list(metric_keys)
    expected_series = {
        str(row["series_name"]).strip()
        for row in rows
        if str(row.get("series_name", "")).strip()
    }
    if not expected_series:
        return [], []

    rows_by_date: dict[str, list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        trade_date = str(row.get("trade_date", "")).strip()
        if trade_date:
            rows_by_date[trade_date].append(row)

    dropped_dates: set[str] = set()
    for trade_date, date_rows in rows_by_date.items():
        series_counter = Counter(str(row["series_name"]).strip() for row in date_rows)
        series_set = {series for series in series_counter if series}
        if series_set != expected_series or any(count != 1 for count in series_counter.values()):
            dropped_dates.add(trade_date)
            continue

        for row in date_rows:
            metrics = row.get("metrics", {})
            if any(metrics.get(metric_key) is None for metric_key in metric_key_list):
                dropped_dates.add(trade_date)
                break

    cleaned_rows = [
        row
        for row in rows
        if str(row.get("trade_date", "")).strip() not in dropped_dates
    ]
    return cleaned_rows, sorted(dropped_dates)


def _deduplicate_generic_rows(
    rows: list[dict[str, object]],
) -> tuple[list[dict[str, object]], int]:
    deduplicated_rows: list[dict[str, object]] = []
    seen_keys: set[tuple[object, ...]] = set()
    duplicate_count = 0

    for row in rows:
        metrics = row.get("metrics", {})
        metric_signature = tuple(sorted(metrics.items()))
        dedupe_key = (
            row.get("dataset_code"),
            row.get("series_name"),
            row.get("trade_date"),
            metric_signature,
        )
        if dedupe_key in seen_keys:
            duplicate_count += 1
            continue
        seen_keys.add(dedupe_key)
        deduplicated_rows.append(row)

    return deduplicated_rows, duplicate_count


def _canonicalize_generic_series_names(
    rows: list[dict[str, object]],
) -> list[dict[str, object]]:
    if not rows:
        return []

    grouped_rows: list[tuple[str, list[dict[str, object]]]] = []
    current_group: list[dict[str, object]] = []
    current_series_name: str | None = None

    for row in rows:
        series_name = str(row.get("series_name", "")).strip()
        if current_series_name is None or series_name == current_series_name:
            current_group.append(row)
            current_series_name = series_name
            continue

        grouped_rows.append((current_series_name, current_group))
        current_group = [row]
        current_series_name = series_name

    if current_group and current_series_name is not None:
        grouped_rows.append((current_series_name, current_group))

    canonical_groups: dict[str, list[tuple[tuple[object, ...], str]]] = defaultdict(list)
    normalized_rows: list[dict[str, object]] = []
    for base_series_name, group_rows in grouped_rows:
        group_signature = tuple(
            (
                row.get("trade_date"),
                tuple(sorted((row.get("metrics") or {}).items())),
            )
            for row in group_rows
        )
        existing_groups = canonical_groups[base_series_name]
        canonical_name = None
        for existing_signature, existing_name in existing_groups:
            if existing_signature == group_signature:
                canonical_name = existing_name
                break

        if canonical_name is None:
            occurrence_index = len(existing_groups) + 1
            canonical_name = (
                base_series_name
                if occurrence_index == 1
                else f"{base_series_name} [{occurrence_index}]"
            )
            existing_groups.append((group_signature, canonical_name))

        for row in group_rows:
            normalized_row = dict(row)
            normalized_row["series_name"] = canonical_name
            normalized_rows.append(normalized_row)

    return normalized_rows


def _finalize_generic_market_payload(
    workbook_path: Path,
    source_sheet: str,
    metric_keys: list[str],
    parsed_rows: list[dict[str, object]],
) -> dict[str, object]:
    dataset_code = _resolve_dataset_code(workbook_path)
    canonical_rows = _canonicalize_generic_series_names(parsed_rows)
    deduplicated_rows, duplicate_row_count = _deduplicate_generic_rows(canonical_rows)
    if dataset_code in SPARSE_GENERIC_DATASET_CODES:
        cleaned_rows = deduplicated_rows
        dropped_dates: list[str] = []
    else:
        cleaned_rows, dropped_dates = _drop_incomplete_dates(
            deduplicated_rows,
            metric_keys=metric_keys,
        )

    return {
        "dataset_code": dataset_code,
        "dataset_name": workbook_path.stem,
        "excel_path": str(workbook_path),
        "source_sheet": source_sheet,
        "metric_keys": metric_keys,
        "source_row_count": len(parsed_rows),
        "deduplicated_row_count": len(deduplicated_rows),
        "duplicate_row_count": duplicate_row_count,
        "clean_row_count": len(cleaned_rows),
        "dropped_date_count": len(dropped_dates),
        "dropped_dates": dropped_dates,
        "series_count": len({row["series_name"] for row in deduplicated_rows}),
        "rows": cleaned_rows,
    }


def _build_generic_record(
    workbook_path: Path,
    *,
    source_sheet: str,
    metric_keys: list[str],
    raw_record: dict[str, str],
) -> dict[str, object] | None:
    series_name = str(raw_record.get("series_name", "")).strip()
    trade_date = _excel_date_to_iso(str(raw_record.get("trade_date", "")))
    if not series_name or not trade_date:
        return None

    metrics = {
        metric_key: _to_float(str(raw_record.get(metric_key, "")))
        for metric_key in metric_keys
    }
    if not any(metric_value is not None for metric_value in metrics.values()):
        return None

    return {
        "dataset_code": _resolve_dataset_code(workbook_path),
        "dataset_name": workbook_path.stem,
        "source_file": workbook_path.name,
        "source_sheet": source_sheet,
        "series_name": series_name,
        "source_label": str(raw_record.get("source_label", "")).strip() or None,
        "frequency_label": str(raw_record.get("frequency_label", "")).strip() or None,
        "unit_label": str(raw_record.get("unit_label", "")).strip() or None,
        "series_id": str(raw_record.get("series_id", "")).strip() or None,
        "trade_date": trade_date,
        "metrics": metrics,
    }


def read_generic_market_workbook(
    excel_path: Path | str,
    sheet_name: str | None = None,
) -> dict[str, object]:
    workbook_path = Path(excel_path)
    if not workbook_path.exists():
        raise FileNotFoundError(f"Excel file not found: {workbook_path}")

    if workbook_path.suffix.lower() == ".csv":
        rows, _encoding = _read_csv_rows(workbook_path)
        if not rows:
            return _finalize_generic_market_payload(workbook_path, workbook_path.stem, [], [])

        mapped_headers = [
            GENERIC_HEADER_MAP.get(_normalize_header_label(value), "")
            for value in rows[0]
        ]
        if "series_name" not in mapped_headers or "trade_date" not in mapped_headers:
            raise ValueError(f"No structured data header found in csv: {workbook_path}")

        metric_keys = [
            header
            for header in mapped_headers
            if header and header not in GENERIC_DIMENSION_KEYS
        ]
        if not metric_keys:
            raise ValueError(f"No metric columns found in csv: {workbook_path}")

        parsed_rows: list[dict[str, object]] = []
        for row in rows[1:]:
            if not any(str(value).strip() for value in row):
                continue
            raw_record: dict[str, str] = {}
            for index, header in enumerate(mapped_headers):
                if not header:
                    continue
                raw_record[header] = str(row[index]).strip() if index < len(row) else ""
            record = _build_generic_record(
                workbook_path,
                source_sheet=workbook_path.stem,
                metric_keys=metric_keys,
                raw_record=raw_record,
            )
            if record is not None:
                parsed_rows.append(record)

        return _finalize_generic_market_payload(
            workbook_path,
            workbook_path.stem,
            metric_keys,
            parsed_rows,
        )

    with zipfile.ZipFile(workbook_path) as zf:
        shared_strings = _load_shared_strings(zf)
        try:
            source_sheet, sheet_rows, header_index, mapped_headers = _find_generic_sheet(
                zf,
                shared_strings,
                sheet_name=sheet_name,
            )
        except ValueError:
            try:
                return _parse_wide_metadata_workbook(
                    workbook_path,
                    zf,
                    shared_strings,
                    sheet_name=sheet_name,
                )
            except ValueError:
                try:
                    return _parse_supply_country_workbook(
                        workbook_path,
                        zf,
                        shared_strings,
                        sheet_name=sheet_name,
                    )
                except ValueError:
                    return _parse_compact_wide_workbook(
                        workbook_path,
                        zf,
                        shared_strings,
                        sheet_name=sheet_name,
                    )

    metric_keys = [
        header
        for header in mapped_headers
        if header and header not in GENERIC_DIMENSION_KEYS
    ]

    parsed_rows: list[dict[str, object]] = []
    for row in sheet_rows[header_index + 1 :]:
        if not any(str(value).strip() for value in row):
            continue

        raw_record: dict[str, str] = {}
        for index, header in enumerate(mapped_headers):
            if not header:
                continue
            raw_record[header] = str(row[index]).strip() if index < len(row) else ""

        record = _build_generic_record(
            workbook_path,
            source_sheet=source_sheet,
            metric_keys=metric_keys,
            raw_record=raw_record,
        )
        if record is not None:
            parsed_rows.append(record)

    return _finalize_generic_market_payload(
        workbook_path,
        source_sheet,
        metric_keys,
        parsed_rows,
    )


def replace_generic_market_rows(
    dataset_code: str,
    rows: Iterable[dict[str, object]],
    db_path: Path | str = DEFAULT_DB_PATH,
) -> int:
    materialized_rows = list(rows)
    init_market_database(db_path)
    with get_connection(db_path) as connection:
        connection.execute(
            """
            DELETE FROM market_generic_daily_metrics
            WHERE dataset_code = ?
            """,
            (dataset_code,),
        )

        payload = []
        for row in materialized_rows:
            metrics = row.get("metrics", {})
            for metric_key, metric_value in metrics.items():
                if metric_value is None:
                    continue
                payload.append(
                    (
                        row["dataset_code"],
                        row["dataset_name"],
                        row["source_file"],
                        row["source_sheet"],
                        row["series_name"],
                        row.get("source_label"),
                        row.get("frequency_label"),
                        row.get("unit_label"),
                        row.get("series_id"),
                        row["trade_date"],
                        metric_key,
                        metric_value,
                    )
                )

        if payload:
            connection.executemany(
                """
                INSERT INTO market_generic_daily_metrics (
                    dataset_code,
                    dataset_name,
                    source_file,
                    source_sheet,
                    series_name,
                    source_label,
                    frequency_label,
                    unit_label,
                    series_id,
                    trade_date,
                    metric_key,
                    metric_value
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(dataset_code, series_name, trade_date, metric_key) DO UPDATE SET
                    dataset_name = excluded.dataset_name,
                    source_file = excluded.source_file,
                    source_sheet = excluded.source_sheet,
                    source_label = excluded.source_label,
                    frequency_label = excluded.frequency_label,
                    unit_label = excluded.unit_label,
                    series_id = excluded.series_id,
                    metric_value = excluded.metric_value
                """,
                payload,
            )
    return len(payload)


def import_generic_market_workbook(
    excel_path: Path | str,
    db_path: Path | str = DEFAULT_DB_PATH,
    sheet_name: str | None = None,
) -> dict[str, object]:
    workbook_path = Path(excel_path)
    if _is_news_extraction_workbook(workbook_path):
        return import_news_extraction_workbook(
            workbook_path,
            db_path=db_path,
            sheet_name=sheet_name,
        )
    if workbook_path.name.lower() in SUPPLY_TRACKING_WORKBOOK_NAMES:
        return import_supply_tracking_workbook(
            workbook_path,
            db_path=db_path,
            sheet_name=sheet_name,
        )
    if _is_port_reference_workbook(workbook_path):
        return import_port_reference_workbook(
            workbook_path,
            db_path=db_path,
            sheet_name=sheet_name,
        )

    workbook = read_generic_market_workbook(excel_path, sheet_name=sheet_name)
    inserted_metric_rows = replace_generic_market_rows(
        workbook["dataset_code"],
        workbook["rows"],
        db_path=db_path,
    )

    return {
        "excel_path": workbook["excel_path"],
        "db_path": str(Path(db_path)),
        "dataset_code": workbook["dataset_code"],
        "dataset_name": workbook["dataset_name"],
        "source_sheet": workbook["source_sheet"],
        "series_count": workbook["series_count"],
        "source_row_count": workbook["source_row_count"],
        "deduplicated_row_count": workbook["deduplicated_row_count"],
        "duplicate_row_count": workbook["duplicate_row_count"],
        "clean_row_count": workbook["clean_row_count"],
        "metric_row_count": inserted_metric_rows,
        "dropped_date_count": workbook["dropped_date_count"],
        "dropped_dates": workbook["dropped_dates"],
        "metric_keys": workbook["metric_keys"],
    }


def import_all_market_data(
    db_path: Path | str = DEFAULT_DB_PATH,
    sample_path: Path | str | None = None,
    workbook_dir: Path | str = DEFAULT_WORKBOOK_DIR,
) -> dict[str, object]:
    summaries: list[dict[str, object]] = []
    workbook_root = Path(workbook_dir)
    sample_file = resolve_sample_xlsx(sample_path)
    if sample_file.exists():
        summaries.append(
            {
                "kind": "sample",
                "summary": import_market_data(sample_file, db_path=db_path),
            }
        )

    macro_split_summaries = split_macro_indicator_csv(
        output_dir=workbook_root,
    )
    for split_summary in macro_split_summaries:
        generated_path = Path(split_summary["output_path"])
        try:
            summaries.append(
                {
                    "kind": "generic",
                    "summary": import_generic_market_workbook(generated_path, db_path=db_path),
                }
            )
        except Exception as error:
            summaries.append(
                {
                    "kind": "error",
                    "summary": {
                        "excel_path": str(generated_path),
                        "error": str(error),
                    },
                }
            )

    for workbook_path in list_additional_market_workbooks(workbook_root):
        try:
            summaries.append(
                {
                    "kind": "generic",
                    "summary": import_generic_market_workbook(workbook_path, db_path=db_path),
                }
            )
        except Exception as error:
            summaries.append(
                {
                    "kind": "error",
                    "summary": {
                        "excel_path": str(workbook_path),
                        "error": str(error),
                    },
                }
            )

    return {
        "db_path": str(Path(db_path)),
        "summaries": summaries,
    }


def upsert_market_rows(
    rows: Iterable[dict[str, object]],
    db_path: Path | str = DEFAULT_DB_PATH,
) -> int:
    payload = [
        (
            row["symbol"],
            row["trade_date"],
            row["change"],
            row["change_pct"],
            row["open_price"],
            row["high_price"],
            row["low_price"],
            row["close_price"],
            row["volume"],
        )
        for row in rows
    ]
    if not payload:
        return 0

    init_market_database(db_path)
    with get_connection(db_path) as connection:
        connection.executemany(
            """
            INSERT INTO market_daily_prices (
                symbol,
                trade_date,
                change,
                change_pct,
                open_price,
                high_price,
                low_price,
                close_price,
                volume
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(symbol, trade_date) DO UPDATE SET
                change = excluded.change,
                change_pct = excluded.change_pct,
                open_price = excluded.open_price,
                high_price = excluded.high_price,
                low_price = excluded.low_price,
                close_price = excluded.close_price,
                volume = excluded.volume
            """,
            payload,
        )
    return len(payload)


def delete_market_rows_by_keys(
    rows: Iterable[dict[str, object]],
    db_path: Path | str = DEFAULT_DB_PATH,
) -> int:
    payload = []
    for row in rows:
        symbol = row.get("symbol")
        trade_date = row.get("trade_date")
        if symbol and trade_date:
            payload.append((symbol, trade_date))

    if not payload:
        return 0

    init_market_database(db_path)
    with get_connection(db_path) as connection:
        cursor = connection.executemany(
            """
            DELETE FROM market_daily_prices
            WHERE symbol = ? AND trade_date = ?
            """,
            payload,
        )
        return cursor.rowcount if cursor.rowcount is not None else 0


def delete_rows_with_missing_open(
    db_path: Path | str = DEFAULT_DB_PATH,
) -> int:
    init_market_database(db_path)
    with get_connection(db_path) as connection:
        cursor = connection.execute(
            """
            DELETE FROM market_daily_prices
            WHERE open_price IS NULL
            """
        )
        return cursor.rowcount if cursor.rowcount is not None else 0


def import_market_data(
    excel_path: Path | str,
    db_path: Path | str = DEFAULT_DB_PATH,
    sheet_name: str | None = None,
) -> dict[str, object]:
    rows = read_market_rows_from_xlsx(excel_path, sheet_name=sheet_name)
    dropped_rows = [row for row in rows if row.get("open_price") is None]
    valid_rows = [row for row in rows if row.get("open_price") is not None]

    deleted_missing_open_keys = delete_market_rows_by_keys(dropped_rows, db_path=db_path)
    inserted = upsert_market_rows(valid_rows, db_path=db_path)
    deleted_missing_open_legacy = delete_rows_with_missing_open(db_path=db_path)

    by_symbol: dict[str, int] = defaultdict(int)
    for row in valid_rows:
        by_symbol[str(row["symbol"])] += 1

    return {
        "excel_path": str(Path(excel_path)),
        "db_path": str(Path(db_path)),
        "row_count": inserted,
        "dropped_missing_open": len(dropped_rows),
        "deleted_missing_open": deleted_missing_open_keys + deleted_missing_open_legacy,
        "symbols": dict(sorted(by_symbol.items())),
    }


def count_market_rows(db_path: Path | str = DEFAULT_DB_PATH) -> int:
    init_market_database(db_path)
    with get_connection(db_path) as connection:
        row = connection.execute("SELECT COUNT(*) AS count FROM market_daily_prices").fetchone()
    return int(row["count"]) if row else 0


def bootstrap_sample_database(
    sample_path: Path | str | None = None,
    db_path: Path | str = DEFAULT_DB_PATH,
) -> dict[str, object] | None:
    sample_file = resolve_sample_xlsx(sample_path)
    if count_market_rows(db_path) > 0:
        return None
    if not sample_file.exists():
        return None
    return import_market_data(sample_file, db_path=db_path)


def get_latest_snapshots(
    symbols: Iterable[str] | None = None,
    db_path: Path | str = DEFAULT_DB_PATH,
) -> list[dict[str, object]]:
    init_market_database(db_path)
    normalized_symbols = [normalize_symbol(symbol) for symbol in symbols or ("WTI", "Brent")]
    normalized_symbols = [symbol for symbol in normalized_symbols if symbol]
    placeholders = ",".join("?" for _ in normalized_symbols)

    query = f"""
        SELECT price.symbol,
               price.trade_date,
               price.change,
               price.change_pct,
               price.open_price,
               price.high_price,
               price.low_price,
               price.close_price,
               price.volume
        FROM market_daily_prices AS price
        INNER JOIN (
            SELECT symbol, MAX(trade_date) AS max_trade_date
            FROM market_daily_prices
            WHERE symbol IN ({placeholders})
            GROUP BY symbol
        ) AS latest
        ON latest.symbol = price.symbol
       AND latest.max_trade_date = price.trade_date
        ORDER BY CASE price.symbol
            WHEN 'WTI' THEN 1
            WHEN 'Brent' THEN 2
            ELSE 99
        END
    """

    with get_connection(db_path) as connection:
        rows = connection.execute(query, normalized_symbols).fetchall()

    return [
        {
            "symbol": row["symbol"],
            "label": row["symbol"],
            "trade_date": row["trade_date"],
            "change": row["change"],
            "change_pct": row["change_pct"],
            "open_price": row["open_price"],
            "high_price": row["high_price"],
            "low_price": row["low_price"],
            "close_price": row["close_price"],
            "volume": row["volume"],
        }
        for row in rows
    ]


def get_price_series(
    symbol: str,
    db_path: Path | str = DEFAULT_DB_PATH,
    *,
    limit: int | None = None,
    date_from: str | None = None,
    date_to: str | None = None,
) -> list[dict[str, object]]:
    normalized_symbol = normalize_symbol(symbol)
    if not normalized_symbol:
        return []

    init_market_database(db_path)
    params: list[object] = [normalized_symbol]
    where_clauses = ["symbol = ?"]
    if date_from:
        where_clauses.append("trade_date >= ?")
        params.append(date_from)
    if date_to:
        where_clauses.append("trade_date <= ?")
        params.append(date_to)

    where_sql = " AND ".join(where_clauses)
    limit_sql = ""
    if limit and limit > 0:
        limit_sql = "LIMIT ?"
        params.append(int(limit))

    query = f"""
        SELECT *
        FROM (
            SELECT symbol,
                   trade_date,
                   change,
                   change_pct,
                   open_price,
                   high_price,
                   low_price,
                   close_price,
                   volume
            FROM market_daily_prices
            WHERE {where_sql}
            ORDER BY trade_date DESC
            {limit_sql}
        ) AS recent
        ORDER BY trade_date ASC
    """

    with get_connection(db_path) as connection:
        rows = connection.execute(query, params).fetchall()

    return [
        {
            "symbol": row["symbol"],
            "trade_date": row["trade_date"],
            "change": row["change"],
            "change_pct": row["change_pct"],
            "open_price": row["open_price"],
            "high_price": row["high_price"],
            "low_price": row["low_price"],
            "close_price": row["close_price"],
            "volume": row["volume"],
        }
        for row in rows
    ]


def get_generic_primary_metric(dataset_code: str) -> str:
    return GENERIC_PRIMARY_METRIC.get(str(dataset_code or "").strip(), "price")


def get_generic_latest_metrics(
    dataset_code: str,
    db_path: Path | str = DEFAULT_DB_PATH,
    *,
    metric_key: str | None = None,
    series_names: Iterable[str] | None = None,
) -> list[dict[str, object]]:
    normalized_dataset_code = str(dataset_code or "").strip()
    if not normalized_dataset_code:
        return []

    primary_metric = str(metric_key or get_generic_primary_metric(normalized_dataset_code)).strip()
    requested_series_names = [
        str(series_name).strip()
        for series_name in (series_names or [])
        if str(series_name).strip()
    ]

    init_market_database(db_path)
    params: list[object] = [normalized_dataset_code]
    where_clauses = ["dataset_code = ?"]
    if requested_series_names:
        placeholders = ",".join("?" for _ in requested_series_names)
        where_clauses.append(f"series_name IN ({placeholders})")
        params.extend(requested_series_names)

    where_sql = " AND ".join(where_clauses)
    query = f"""
        SELECT dataset_code,
               dataset_name,
               source_file,
               source_sheet,
               series_name,
               source_label,
               frequency_label,
               unit_label,
               series_id,
               trade_date,
               metric_key,
               metric_value
        FROM market_generic_daily_metrics
        WHERE {where_sql}
        ORDER BY series_name ASC, trade_date ASC
    """

    with get_connection(db_path) as connection:
        rows = connection.execute(query, params).fetchall()

    grouped_rows: dict[str, dict[str, dict[str, object]]] = defaultdict(dict)
    dataset_name = None
    series_meta: dict[str, dict[str, object]] = {}
    for row in rows:
        dataset_name = dataset_name or row["dataset_name"]
        series_name = row["series_name"]
        meta = series_meta.setdefault(
            series_name,
            {
                "source_file": row["source_file"],
                "source_sheet": row["source_sheet"],
                "source_label": row["source_label"],
                "frequency_label": row["frequency_label"],
                "unit_label": row["unit_label"],
                "series_id": row["series_id"],
            },
        )
        for key in (
            "source_file",
            "source_sheet",
            "source_label",
            "frequency_label",
            "unit_label",
            "series_id",
        ):
            meta[key] = meta.get(key) or row[key]
        trade_date = row["trade_date"]
        grouped_rows[series_name].setdefault(trade_date, {})
        grouped_rows[series_name][trade_date][row["metric_key"]] = row["metric_value"]

    if requested_series_names:
        ordered_series_names = requested_series_names
    else:
        ordered_series_names = sorted(grouped_rows)

    items: list[dict[str, object]] = []
    for series_name in ordered_series_names:
        date_map = grouped_rows.get(series_name, {})
        if not date_map:
            continue

        ordered_dates = sorted(date_map)
        latest_trade_date = ordered_dates[-1]
        latest_metrics = date_map[latest_trade_date]
        latest_value = latest_metrics.get(primary_metric)
        if latest_value is None:
            continue

        previous_value = None
        for trade_date in reversed(ordered_dates[:-1]):
            candidate = date_map[trade_date].get(primary_metric)
            if candidate is not None:
                previous_value = candidate
                break

        change = None
        change_pct = None
        if previous_value is not None:
            change = float(latest_value) - float(previous_value)
            change_pct = (change / float(previous_value) * 100) if previous_value else None

        items.append(
            {
                "dataset_code": normalized_dataset_code,
                "dataset_name": dataset_name,
                "source_file": series_meta.get(series_name, {}).get("source_file"),
                "source_sheet": series_meta.get(series_name, {}).get("source_sheet"),
                "source_label": series_meta.get(series_name, {}).get("source_label"),
                "frequency_label": series_meta.get(series_name, {}).get("frequency_label"),
                "unit_label": series_meta.get(series_name, {}).get("unit_label"),
                "series_id": series_meta.get(series_name, {}).get("series_id"),
                "series_name": series_name,
                "label": series_name,
                "trade_date": latest_trade_date,
                "metric_key": primary_metric,
                "value": float(latest_value),
                "change": round(change, 6) if change is not None else None,
                "change_pct": round(change_pct, 6) if change_pct is not None else None,
                "metrics": {
                    metric_name: float(metric_value)
                    for metric_name, metric_value in latest_metrics.items()
                    if metric_value is not None
                },
            }
        )

    return items


def get_generic_metric_series(
    dataset_code: str,
    series_name: str,
    db_path: Path | str = DEFAULT_DB_PATH,
    *,
    metric_key: str | None = None,
    limit: int | None = None,
    date_from: str | None = None,
    date_to: str | None = None,
) -> list[dict[str, object]]:
    normalized_dataset_code = str(dataset_code or "").strip()
    normalized_series_name = str(series_name or "").strip()
    if not normalized_dataset_code or not normalized_series_name:
        return []

    primary_metric = str(metric_key or get_generic_primary_metric(normalized_dataset_code)).strip()

    init_market_database(db_path)
    params: list[object] = [normalized_dataset_code, normalized_series_name, primary_metric]
    where_clauses = [
        "dataset_code = ?",
        "series_name = ?",
        "metric_key = ?",
    ]
    if date_from:
        where_clauses.append("trade_date >= ?")
        params.append(date_from)
    if date_to:
        where_clauses.append("trade_date <= ?")
        params.append(date_to)

    where_sql = " AND ".join(where_clauses)
    limit_sql = ""
    if limit and limit > 0:
        limit_sql = "LIMIT ?"
        params.append(int(limit))

    query = f"""
        SELECT *
        FROM (
            SELECT dataset_code,
                   dataset_name,
                   source_file,
                   source_sheet,
                   series_name,
                   source_label,
                   frequency_label,
                   unit_label,
                   series_id,
                   trade_date,
                   metric_key,
                   metric_value
            FROM market_generic_daily_metrics
            WHERE {where_sql}
            ORDER BY trade_date DESC
            {limit_sql}
        ) AS recent
        ORDER BY trade_date ASC
    """

    with get_connection(db_path) as connection:
        rows = connection.execute(query, params).fetchall()

    return [
        {
            "dataset_code": row["dataset_code"],
            "dataset_name": row["dataset_name"],
            "source_file": row["source_file"],
            "source_sheet": row["source_sheet"],
            "source_label": row["source_label"],
            "frequency_label": row["frequency_label"],
            "unit_label": row["unit_label"],
            "series_id": row["series_id"],
            "series_name": row["series_name"],
            "trade_date": row["trade_date"],
            "metric_key": row["metric_key"],
            "value": row["metric_value"],
        }
        for row in rows
    ]


def _days_in_trade_month(trade_date: str) -> int:
    parsed_date = datetime.strptime(trade_date, "%Y-%m-%d")
    return calendar.monthrange(parsed_date.year, parsed_date.month)[1]


def _normalize_supply_dashboard_value(series_name: str, trade_date: str, raw_value: float) -> float:
    series_key = str(series_name or "").strip()
    if series_key in {
        SUPPLY_DASHBOARD_SERIES["globalSupply"],
        SUPPLY_DASHBOARD_SERIES["globalDemand"],
        SUPPLY_DASHBOARD_SERIES["balance"],
    }:
        return raw_value / 1000.0 / float(_days_in_trade_month(trade_date))
    if series_key == SUPPLY_DASHBOARD_SERIES["floatingStorage"]:
        return raw_value / 1000.0
    return raw_value


def _fallback_country_code(country_name: str) -> str:
    normalized = re.sub(r"[^A-Za-z]+", " ", str(country_name or "")).strip()
    if not normalized:
        return "NA"
    tokens = [token for token in normalized.split() if token]
    if not tokens:
        return "NA"
    if len(tokens) == 1:
        return tokens[0][:2].upper()
    return "".join(token[0].upper() for token in tokens[:3])


def _normalize_supply_tracking_header_label(raw_value: str) -> str:
    return re.sub(r"[\s_]+", "", str(raw_value or "").strip().lower())


def _normalize_ascii_text(raw_value: str | None) -> str:
    normalized = unicodedata.normalize("NFKD", str(raw_value or ""))
    normalized = normalized.replace("’", "'").replace("‘", "'")
    return normalized.encode("ascii", "ignore").decode("ascii")


def _normalize_supply_tracking_token(raw_value: str | None) -> str:
    value = _normalize_ascii_text(raw_value).lower()
    value = value.replace("&", " and ")
    value = value.replace("(", " ").replace(")", " ")
    value = value.replace(">", " ")
    value = re.sub(r"[^a-z0-9/;,]+", " ", value)
    value = re.sub(r"\s+", " ", value).strip()
    return value


def _strip_supply_tracking_navigation_terms(raw_value: str | None) -> str:
    value = _normalize_supply_tracking_token(raw_value)
    value = re.sub(r"\banch(?:orage)?\.?\b", " ", value)
    value = re.sub(r"\bterm(?:inal)?s?\.?\b", " ", value)
    value = re.sub(r"\bport of\b", " ", value)
    value = re.sub(r"\s+", " ", value).strip(" ,;/")
    return value


def _normalize_supply_tracking_country(raw_value: str | None) -> str:
    value = _normalize_supply_tracking_token(raw_value)
    value = re.sub(r"[^a-z0-9]+", " ", value).strip()
    replacements = {
        "usa": "united states",
        "united states usa": "united states",
        "uae": "united arab emirates",
        "united arab emirates uae": "united arab emirates",
    }
    return replacements.get(value, value)


def _extract_supply_tracking_country(raw_value: str | None) -> str:
    text = str(raw_value or "").strip()
    if "," not in text:
        return ""
    return _normalize_supply_tracking_country(text.rsplit(",", 1)[-1])


def _is_supply_tracking_placeholder_value(raw_value: str | None) -> bool:
    normalized = re.sub(r"[^a-z0-9]+", " ", _normalize_supply_tracking_token(raw_value)).strip()
    return not normalized or normalized in SUPPLY_TRACKING_PLACEHOLDER_TOKENS


def _normalize_supply_tracking_place_key(raw_value: str | None) -> str:
    if _is_supply_tracking_placeholder_value(raw_value):
        return ""
    return re.sub(r"[^a-z0-9]+", " ", _normalize_supply_tracking_token(raw_value)).strip()


def _has_same_supply_tracking_route_endpoints(raw_record: dict[str, str]) -> bool:
    place_keys = [
        _normalize_supply_tracking_place_key(raw_record.get("origin_name")),
        _normalize_supply_tracking_place_key(raw_record.get("previous_port_name")),
        _normalize_supply_tracking_place_key(raw_record.get("destination_name")),
    ]
    return bool(place_keys[0] and place_keys[1] and place_keys[2]) and len(set(place_keys)) == 1


def _parse_signed_coordinate_component(
    raw_value: str,
    positive_direction: str,
    negative_direction: str,
) -> float | None:
    match = re.match(
        r"^\s*([+-]?\d+(?:\.\d+)?)\s*°?\s*([A-Za-z])?\s*$",
        str(raw_value or "").strip(),
    )
    if not match:
        return None

    magnitude = float(match.group(1))
    direction = str(match.group(2) or "").upper()
    if direction == negative_direction:
        return -abs(magnitude)
    if direction == positive_direction:
        return abs(magnitude)
    if direction:
        return None
    return magnitude


def _parse_supply_tracking_position(
    raw_value: str | None,
) -> tuple[float, float] | None:
    text = str(raw_value or "").strip()
    if not text:
        return None

    parts = [part.strip() for part in text.split(",")]
    if len(parts) >= 2:
        latitude = _parse_signed_coordinate_component(parts[0], "N", "S")
        longitude = _parse_signed_coordinate_component(parts[1], "E", "W")
        if latitude is not None and longitude is not None:
            return latitude, longitude

    numeric_parts = re.findall(r"[-+]?\d+(?:\.\d+)?", text)
    if len(numeric_parts) >= 2:
        latitude = float(numeric_parts[0])
        longitude = float(numeric_parts[1])
        return latitude, longitude
    return None


def _read_supply_tracking_sheet_rows(
    workbook_path: Path,
    sheet_name: str | None = None,
) -> tuple[str, list[list[str]]]:
    if workbook_path.suffix.lower() == ".csv":
        rows, _encoding = _read_csv_rows(workbook_path)
        return workbook_path.stem, rows

    with zipfile.ZipFile(workbook_path) as zf:
        shared_strings = _load_shared_strings(zf)
        sheet_targets = _sheet_targets(zf)
        if sheet_name:
            sheet_targets = [item for item in sheet_targets if item[0] == sheet_name]
            if not sheet_targets:
                raise ValueError(f"Sheet not found: {sheet_name}")
        if not sheet_targets:
            raise ValueError(f"No worksheet found in workbook: {workbook_path}")
        source_sheet, target_path = sheet_targets[0]
        return source_sheet, _read_sheet_rows(zf, target_path, shared_strings)


def _find_supply_tracking_header(
    rows: list[list[str]],
) -> tuple[int, list[str]]:
    for row_index, row in enumerate(rows[:10]):
        mapped_headers = [
            SUPPLY_TRACKING_NORMALIZED_HEADER_MAP.get(
                _normalize_supply_tracking_header_label(value),
                "",
            )
            for value in row
        ]
        if all(field in mapped_headers for field in SUPPLY_TRACKING_REQUIRED_FIELDS):
            return row_index, mapped_headers
    raise ValueError("No structured tanker tracking header found")


def read_supply_tracking_workbook(
    excel_path: Path | str,
    sheet_name: str | None = None,
) -> dict[str, object]:
    workbook_path = Path(excel_path)
    source_sheet, rows = _read_supply_tracking_sheet_rows(workbook_path, sheet_name=sheet_name)
    if not rows:
        return {
            "excel_path": str(workbook_path),
            "dataset_code": SUPPLY_TRACKING_DATASET_CODE,
            "dataset_name": workbook_path.stem,
            "source_sheet": source_sheet,
            "source_row_count": 0,
            "clean_row_count": 0,
            "dropped_row_count": 0,
            "dropped_rows": [],
            "rows": [],
        }

    header_index, mapped_headers = _find_supply_tracking_header(rows)
    parsed_rows: list[dict[str, object]] = []
    dropped_rows: list[int] = []

    for source_row_number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        if not any(str(value).strip() for value in row):
            continue

        raw_record: dict[str, str] = {}
        for index, mapped_header in enumerate(mapped_headers):
            if not mapped_header:
                continue
            raw_record[mapped_header] = str(row[index]).strip() if index < len(row) else ""

        if _has_same_supply_tracking_route_endpoints(raw_record):
            dropped_rows.append(source_row_number)
            continue

        vessel_name = str(raw_record.get("vessel_name") or "").strip()
        current_position_raw = str(raw_record.get("current_position_raw") or "").strip()
        if not vessel_name and not current_position_raw:
            continue

        current_coords = _parse_supply_tracking_position(current_position_raw)
        if current_coords is None:
            dropped_rows.append(source_row_number)
            continue

        parsed_rows.append(
            {
                "dataset_code": SUPPLY_TRACKING_DATASET_CODE,
                "dataset_name": workbook_path.stem,
                "source_file": workbook_path.name,
                "source_sheet": source_sheet,
                "source_row_number": source_row_number,
                "vessel_name": vessel_name or f"Unknown Vessel {source_row_number}",
                "imo": str(raw_record.get("imo") or "").strip() or None,
                "mmsi": str(raw_record.get("mmsi") or "").strip() or None,
                "vessel_type": str(raw_record.get("vessel_type") or "").strip() or "Unknown",
                "dwt": _to_float(str(raw_record.get("dwt") or "").strip()),
                "current_position_raw": current_position_raw or None,
                "current_latitude": current_coords[0],
                "current_longitude": current_coords[1],
                "position_label": current_position_raw or None,
                "origin_name": str(raw_record.get("origin_name") or "").strip() or None,
                "previous_port_name": str(raw_record.get("previous_port_name") or "").strip() or None,
                "destination_name": str(raw_record.get("destination_name") or "").strip() or None,
                "speed": None,
                "heading": None,
                "cargo_status": "unknown",
                "eta": None,
                "source_label": SUPPLY_TRACKING_SOURCE_LABEL,
            }
        )

    return {
        "excel_path": str(workbook_path),
        "dataset_code": SUPPLY_TRACKING_DATASET_CODE,
        "dataset_name": workbook_path.stem,
        "source_sheet": source_sheet,
        "source_row_count": max(len(rows) - header_index - 1, 0),
        "clean_row_count": len(parsed_rows),
        "dropped_row_count": len(dropped_rows),
        "dropped_rows": dropped_rows,
        "rows": parsed_rows,
    }


def _split_port_aliases(raw_value: str | None) -> list[str]:
    normalized = _normalize_ascii_text(raw_value)
    if not normalized.strip():
        return []
    values = [normalized]
    values.extend(re.split(r"[;/]", normalized))
    aliases: list[str] = []
    for value in values:
        for candidate in (
            _normalize_supply_tracking_token(value),
            _strip_supply_tracking_navigation_terms(value),
        ):
            candidate = re.sub(r"[^a-z0-9]+", " ", candidate).strip()
            if candidate and candidate not in aliases:
                aliases.append(candidate)
    return aliases


def _build_supply_tracking_port_lookup(
    db_path: Path | str = DEFAULT_DB_PATH,
) -> dict[str, object]:
    init_market_database(db_path)
    with get_connection(db_path) as connection:
        rows = connection.execute(
            """
            SELECT main_port_name,
                   alternate_port_name,
                   country_name,
                   latitude,
                   longitude
            FROM port_reference_data
            WHERE dataset_code = ?
            """,
            (SUPPLY_PORT_DATASET_CODE,),
        ).fetchall()

    ports: list[dict[str, object]] = []
    by_key: dict[str, list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        alias_keys = _split_port_aliases(row["main_port_name"])
        alias_keys.extend(_split_port_aliases(row["alternate_port_name"]))
        deduped_alias_keys = list(dict.fromkeys(alias_keys))
        port = {
            "name": str(row["main_port_name"]).strip(),
            "country": str(row["country_name"] or "").strip() or None,
            "country_key": _normalize_supply_tracking_country(row["country_name"]),
            "latitude": float(row["latitude"]),
            "longitude": float(row["longitude"]),
            "alias_keys": deduped_alias_keys,
        }
        ports.append(port)
        for key in deduped_alias_keys:
            by_key[key].append(port)

    return {"ports": ports, "by_key": by_key}


def _select_supply_tracking_exact_port(
    matches: list[dict[str, object]],
    country_key: str,
) -> dict[str, object] | None:
    if not matches:
        return None
    if country_key:
        country_matches = [port for port in matches if port["country_key"] == country_key]
        if country_matches:
            return country_matches[0]
    return matches[0]


def _build_supply_tracking_place_candidates(raw_value: str | None) -> list[str]:
    text = str(raw_value or "").strip()
    if not text:
        return []

    normalized_full = re.sub(r"[^a-z0-9]+", " ", _normalize_supply_tracking_token(text)).strip()
    base_text = text.rsplit(",", 1)[0].strip() if "," in text else text
    variants = [text, base_text]
    variants.extend(re.split(r"[;/]", base_text))
    candidates: list[str] = []

    alias_target = SUPPLY_TRACKING_PORT_ALIAS_MAP.get(normalized_full)
    if alias_target:
        variants.insert(0, alias_target)

    for value in variants:
        for candidate in (
            _normalize_supply_tracking_token(value),
            _strip_supply_tracking_navigation_terms(value),
        ):
            normalized = re.sub(r"[^a-z0-9]+", " ", candidate).strip()
            if normalized and normalized not in candidates:
                candidates.append(normalized)
    return candidates


def _match_supply_tracking_port(
    raw_value: str | None,
    port_lookup: dict[str, object],
) -> dict[str, object]:
    if _is_supply_tracking_placeholder_value(raw_value):
        return {
            "name": str(raw_value or "").strip() or None,
            "latitude": None,
            "longitude": None,
            "match_method": None,
            "matched_port_name": None,
        }

    candidates = _build_supply_tracking_place_candidates(raw_value)
    country_key = _extract_supply_tracking_country(raw_value)
    by_key = port_lookup["by_key"]

    for candidate in candidates:
        exact_match = _select_supply_tracking_exact_port(list(by_key.get(candidate, [])), country_key)
        if exact_match is not None:
            return {
                "name": str(raw_value or "").strip() or None,
                "latitude": float(exact_match["latitude"]),
                "longitude": float(exact_match["longitude"]),
                "match_method": "port_reference_exact",
                "matched_port_name": str(exact_match["name"]),
            }

    scored_matches: list[tuple[int, dict[str, object]]] = []
    for port in port_lookup["ports"]:
        if country_key and port["country_key"] and port["country_key"] != country_key:
            continue
        best_score = 0
        for candidate in candidates:
            if len(candidate) < 4:
                continue
            for alias_key in port["alias_keys"]:
                if candidate == alias_key:
                    best_score = max(best_score, 1000 + len(candidate))
                elif candidate in alias_key or alias_key in candidate:
                    best_score = max(best_score, min(len(candidate), len(alias_key)))
        if best_score > 0:
            scored_matches.append((best_score, port))

    if scored_matches:
        scored_matches.sort(
            key=lambda item: (
                item[0],
                1 if country_key and item[1]["country_key"] == country_key else 0,
                len(item[1]["name"]),
            ),
            reverse=True,
        )
        best_port = scored_matches[0][1]
        return {
            "name": str(raw_value or "").strip() or None,
            "latitude": float(best_port["latitude"]),
            "longitude": float(best_port["longitude"]),
            "match_method": "port_reference_fuzzy",
            "matched_port_name": str(best_port["name"]),
        }

    return {
        "name": str(raw_value or "").strip() or None,
        "latitude": None,
        "longitude": None,
        "match_method": None,
        "matched_port_name": None,
    }


def _enrich_supply_tracking_rows(
    rows: list[dict[str, object]],
    db_path: Path | str = DEFAULT_DB_PATH,
) -> list[dict[str, object]]:
    port_lookup = _build_supply_tracking_port_lookup(db_path)
    enriched_rows: list[dict[str, object]] = []

    for row in rows:
        current_latitude = float(row["current_latitude"])
        current_longitude = float(row["current_longitude"])
        origin_match = _match_supply_tracking_port(row.get("origin_name"), port_lookup)
        previous_match = _match_supply_tracking_port(row.get("previous_port_name"), port_lookup)
        destination_match = _match_supply_tracking_port(row.get("destination_name"), port_lookup)

        if origin_match["latitude"] is None:
            origin_match["latitude"] = current_latitude
            origin_match["longitude"] = current_longitude
            origin_match["match_method"] = "fallback_current_position"

        if previous_match["latitude"] is None:
            previous_match["latitude"] = current_latitude
            previous_match["longitude"] = current_longitude
            previous_match["match_method"] = "fallback_current_position"

        enriched_rows.append(
            {
                **row,
                "origin_latitude": origin_match["latitude"],
                "origin_longitude": origin_match["longitude"],
                "origin_match_method": origin_match["match_method"],
                "origin_matched_port_name": origin_match["matched_port_name"],
                "previous_port_latitude": previous_match["latitude"],
                "previous_port_longitude": previous_match["longitude"],
                "previous_port_match_method": previous_match["match_method"],
                "previous_port_matched_port_name": previous_match["matched_port_name"],
                "destination_latitude": destination_match["latitude"],
                "destination_longitude": destination_match["longitude"],
                "destination_match_method": destination_match["match_method"],
                "destination_matched_port_name": destination_match["matched_port_name"],
            }
        )

    return enriched_rows


def replace_supply_tracking_rows(
    dataset_code: str,
    rows: Iterable[dict[str, object]],
    db_path: Path | str = DEFAULT_DB_PATH,
) -> int:
    materialized_rows = list(rows)
    init_market_database(db_path)
    with get_connection(db_path) as connection:
        connection.execute(
            """
            DELETE FROM supply_tanker_tracking
            WHERE dataset_code = ?
            """,
            (dataset_code,),
        )
        if not materialized_rows:
            return 0

        payload = [
            (
                row["dataset_code"],
                row["dataset_name"],
                row["source_file"],
                row["source_sheet"],
                row["source_row_number"],
                row["vessel_name"],
                row.get("imo"),
                row.get("mmsi"),
                row["vessel_type"],
                row.get("dwt"),
                row.get("current_position_raw"),
                row["current_latitude"],
                row["current_longitude"],
                row.get("position_label"),
                row.get("origin_name"),
                row.get("origin_latitude"),
                row.get("origin_longitude"),
                row.get("origin_match_method"),
                row.get("origin_matched_port_name"),
                row.get("previous_port_name"),
                row.get("previous_port_latitude"),
                row.get("previous_port_longitude"),
                row.get("previous_port_match_method"),
                row.get("previous_port_matched_port_name"),
                row.get("destination_name"),
                row.get("destination_latitude"),
                row.get("destination_longitude"),
                row.get("destination_match_method"),
                row.get("destination_matched_port_name"),
                row.get("speed"),
                row.get("heading"),
                row.get("cargo_status") or "unknown",
                row.get("eta"),
                row.get("source_label"),
            )
            for row in materialized_rows
        ]
        connection.executemany(
            """
            INSERT INTO supply_tanker_tracking (
                dataset_code,
                dataset_name,
                source_file,
                source_sheet,
                source_row_number,
                vessel_name,
                imo,
                mmsi,
                vessel_type,
                dwt,
                current_position_raw,
                current_latitude,
                current_longitude,
                position_label,
                origin_name,
                origin_latitude,
                origin_longitude,
                origin_match_method,
                origin_matched_port_name,
                previous_port_name,
                previous_port_latitude,
                previous_port_longitude,
                previous_port_match_method,
                previous_port_matched_port_name,
                destination_name,
                destination_latitude,
                destination_longitude,
                destination_match_method,
                destination_matched_port_name,
                speed,
                heading,
                cargo_status,
                eta,
                source_label
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            payload,
        )
    return len(materialized_rows)


def import_supply_tracking_workbook(
    excel_path: Path | str,
    db_path: Path | str = DEFAULT_DB_PATH,
    sheet_name: str | None = None,
) -> dict[str, object]:
    workbook = read_supply_tracking_workbook(excel_path, sheet_name=sheet_name)
    enriched_rows = _enrich_supply_tracking_rows(list(workbook["rows"]), db_path=db_path)
    inserted_rows = replace_supply_tracking_rows(
        workbook["dataset_code"],
        enriched_rows,
        db_path=db_path,
    )
    return {
        "excel_path": workbook["excel_path"],
        "db_path": str(Path(db_path)),
        "dataset_code": workbook["dataset_code"],
        "dataset_name": workbook["dataset_name"],
        "source_sheet": workbook["source_sheet"],
        "source_row_count": workbook["source_row_count"],
        "clean_row_count": workbook["clean_row_count"],
        "dropped_row_count": workbook["dropped_row_count"],
        "dropped_rows": workbook["dropped_rows"],
        "tracking_row_count": inserted_rows,
    }


def count_supply_tracking_rows(
    db_path: Path | str = DEFAULT_DB_PATH,
) -> int:
    init_market_database(db_path)
    with get_connection(db_path) as connection:
        row = connection.execute(
            """
            SELECT COUNT(*) AS count
            FROM supply_tanker_tracking
            WHERE dataset_code = ?
            """,
            (SUPPLY_TRACKING_DATASET_CODE,),
        ).fetchone()
    return int(row["count"]) if row else 0


def bootstrap_supply_tracking_database(
    source_path: Path | str | None = None,
    db_path: Path | str = DEFAULT_DB_PATH,
) -> dict[str, object] | None:
    workbook_path = resolve_supply_tracking_workbook(source_path)
    if count_supply_tracking_rows(db_path) > 0:
        return None
    if not workbook_path.exists():
        return None
    return import_supply_tracking_workbook(workbook_path, db_path=db_path)


def get_supply_tracking_data(
    db_path: Path | str = DEFAULT_DB_PATH,
) -> dict[str, object]:
    init_market_database(db_path)
    with get_connection(db_path) as connection:
        rows = connection.execute(
            """
            SELECT *
            FROM supply_tanker_tracking
            WHERE dataset_code = ?
            ORDER BY vessel_type ASC, COALESCE(dwt, 0) DESC, vessel_name ASC
            """,
            (SUPPLY_TRACKING_DATASET_CODE,),
        ).fetchall()

    items: list[dict[str, object]] = []
    for row in rows:
        items.append(
            {
                "id": str(row["mmsi"] or row["id"]),
                "name": str(row["vessel_name"]),
                "imo": str(row["imo"]).strip() if row["imo"] else None,
                "mmsi": str(row["mmsi"]).strip() if row["mmsi"] else None,
                "type": str(row["vessel_type"]),
                "dwt": float(row["dwt"]) if row["dwt"] is not None else None,
                "latitude": float(row["current_latitude"]),
                "longitude": float(row["current_longitude"]),
                "speed": float(row["speed"]) if row["speed"] is not None else None,
                "heading": float(row["heading"]) if row["heading"] is not None else None,
                "cargoStatus": str(row["cargo_status"] or "unknown"),
                "destination": str(row["destination_name"]).strip() if row["destination_name"] else None,
                "eta": str(row["eta"]).strip() if row["eta"] else None,
                "positionLabel": str(row["position_label"]).strip() if row["position_label"] else None,
                "originPort": {
                    "role": "origin",
                    "name": str(row["origin_name"]).strip() if row["origin_name"] else "",
                    "latitude": float(row["origin_latitude"]) if row["origin_latitude"] is not None else None,
                    "longitude": float(row["origin_longitude"]) if row["origin_longitude"] is not None else None,
                    "matchMethod": str(row["origin_match_method"]).strip() if row["origin_match_method"] else None,
                    "matchedPortName": str(row["origin_matched_port_name"]).strip()
                    if row["origin_matched_port_name"]
                    else None,
                }
                if row["origin_name"]
                else None,
                "previousPort": {
                    "role": "previous",
                    "name": str(row["previous_port_name"]).strip() if row["previous_port_name"] else "",
                    "latitude": float(row["previous_port_latitude"])
                    if row["previous_port_latitude"] is not None
                    else None,
                    "longitude": float(row["previous_port_longitude"])
                    if row["previous_port_longitude"] is not None
                    else None,
                    "matchMethod": str(row["previous_port_match_method"]).strip()
                    if row["previous_port_match_method"]
                    else None,
                    "matchedPortName": str(row["previous_port_matched_port_name"]).strip()
                    if row["previous_port_matched_port_name"]
                    else None,
                }
                if row["previous_port_name"]
                else None,
                "destinationPort": {
                    "role": "destination",
                    "name": str(row["destination_name"]).strip() if row["destination_name"] else "",
                    "latitude": float(row["destination_latitude"])
                    if row["destination_latitude"] is not None
                    else None,
                    "longitude": float(row["destination_longitude"])
                    if row["destination_longitude"] is not None
                    else None,
                    "matchMethod": str(row["destination_match_method"]).strip()
                    if row["destination_match_method"]
                    else None,
                    "matchedPortName": str(row["destination_matched_port_name"]).strip()
                    if row["destination_matched_port_name"]
                    else None,
                }
                if row["destination_name"]
                else None,
                "sourceLabel": str(row["source_label"]).strip() if row["source_label"] else None,
            }
        )

    return {
        "dataset_code": SUPPLY_TRACKING_DATASET_CODE,
        "total": len(items),
        "items": items,
    }


def _normalize_supply_country_value(
    metric_key: str,
    trade_date: str,
    raw_value: float,
) -> float:
    normalized_metric_key = str(metric_key or "").strip()
    if normalized_metric_key in {"production", "imports", "exports", "products_supplied"}:
        return raw_value / 1000.0 / float(_days_in_trade_month(trade_date))
    if normalized_metric_key == "stock_change":
        return raw_value / 1000.0
    return raw_value


def _compute_percent_change(current_value: float | None, previous_value: float | None) -> float | None:
    if current_value is None or previous_value is None:
        return None
    if abs(previous_value) < 1e-9:
        return 0.0 if abs(current_value) < 1e-9 else None
    return (current_value - previous_value) / abs(previous_value) * 100.0


def get_supply_dashboard_data(
    db_path: Path | str = DEFAULT_DB_PATH,
) -> dict[str, object]:
    series_names = list(SUPPLY_DASHBOARD_SERIES.values())
    init_market_database(db_path)
    placeholders = ",".join("?" for _ in series_names)
    query = f"""
        SELECT series_name,
               trade_date,
               metric_value
        FROM market_generic_daily_metrics
        WHERE dataset_code = ?
          AND metric_key = ?
          AND series_name IN ({placeholders})
        ORDER BY trade_date ASC, series_name ASC
    """
    params: list[object] = [SUPPLY_DASHBOARD_DATASET_CODE, "value", *series_names]
    with get_connection(db_path) as connection:
        rows = connection.execute(query, params).fetchall()

    rows_by_date: dict[str, dict[str, float]] = defaultdict(dict)
    for row in rows:
        series_name = str(row["series_name"]).strip()
        trade_date = str(row["trade_date"]).strip()
        metric_value = row["metric_value"]
        if not trade_date or metric_value is None:
            continue
        rows_by_date[trade_date][series_name] = _normalize_supply_dashboard_value(
            series_name,
            trade_date,
            float(metric_value),
        )

    ordered_dates = [
        trade_date
        for trade_date in sorted(rows_by_date)
        if all(series_name in rows_by_date[trade_date] for series_name in series_names)
    ]
    if not ordered_dates:
        return {
            "dataset_code": SUPPLY_DASHBOARD_DATASET_CODE,
            "as_of": None,
            "metrics": {},
            "trend": [],
        }

    def metric_value_at(trade_date: str, metric_name: str) -> float:
        return float(rows_by_date[trade_date][SUPPLY_DASHBOARD_SERIES[metric_name]])

    latest_date = ordered_dates[-1]
    previous_date = ordered_dates[-2] if len(ordered_dates) > 1 else None

    metrics = {}
    for metric_name, unit in (
        ("globalSupply", "M BBL/D"),
        ("globalDemand", "M BBL/D"),
        ("balance", "M BBL/D"),
        ("floatingStorage", "M BBL"),
    ):
        latest_value = metric_value_at(latest_date, metric_name)
        previous_value = metric_value_at(previous_date, metric_name) if previous_date else None
        change_value = None if previous_value is None else latest_value - previous_value
        change_pct = _compute_percent_change(latest_value, previous_value)
        metrics[metric_name] = {
            "value": round(latest_value, 4),
            "unit": unit,
            "change_value": round(change_value, 4) if change_value is not None else None,
            "change_pct": round(change_pct, 4) if change_pct is not None else None,
        }

    trend = [
        {
            "date": trade_date[:7],
            "supply": round(metric_value_at(trade_date, "globalSupply"), 4),
            "demand": round(metric_value_at(trade_date, "globalDemand"), 4),
            "balance": round(metric_value_at(trade_date, "balance"), 4),
            "floatingStorage": round(metric_value_at(trade_date, "floatingStorage"), 4),
        }
        for trade_date in ordered_dates
    ]

    return {
        "dataset_code": SUPPLY_DASHBOARD_DATASET_CODE,
        "as_of": latest_date,
        "metrics": metrics,
        "trend": trend,
    }


def get_supply_country_details(
    db_path: Path | str = DEFAULT_DB_PATH,
) -> dict[str, object]:
    init_market_database(db_path)
    with get_connection(db_path) as connection:
        latest_row = connection.execute(
            """
            SELECT MAX(trade_date) AS latest_date
            FROM market_generic_daily_metrics
            WHERE dataset_code = ?
            """,
            (SUPPLY_COUNTRY_DATASET_CODE,),
        ).fetchone()
        latest_date = str(latest_row["latest_date"]).strip() if latest_row and latest_row["latest_date"] else None
        if not latest_date:
            return {
                "dataset_code": SUPPLY_COUNTRY_DATASET_CODE,
                "as_of": None,
                "items": [],
            }

        order_rows = connection.execute(
            """
            SELECT series_name, MIN(id) AS first_id
            FROM market_generic_daily_metrics
            WHERE dataset_code = ?
            GROUP BY series_name
            ORDER BY first_id ASC
            """,
            (SUPPLY_COUNTRY_DATASET_CODE,),
        ).fetchall()
        ordered_countries = [str(row["series_name"]).strip() for row in order_rows if str(row["series_name"]).strip()]

        rows = connection.execute(
            """
            SELECT series_name,
                   source_label,
                   trade_date,
                   metric_key,
                   metric_value
            FROM market_generic_daily_metrics
            WHERE dataset_code = ?
              AND trade_date = ?
            ORDER BY id ASC
            """,
            (SUPPLY_COUNTRY_DATASET_CODE, latest_date),
        ).fetchall()

    metrics_by_country: dict[str, dict[str, float]] = defaultdict(dict)
    source_by_country: dict[str, str | None] = {}
    for row in rows:
        country = str(row["series_name"]).strip()
        metric_key = str(row["metric_key"]).strip()
        metric_value = row["metric_value"]
        if not country:
            continue
        source_by_country[country] = source_by_country.get(country) or row["source_label"]
        if metric_value is None:
            continue
        metrics_by_country[country][metric_key] = round(
            _normalize_supply_country_value(metric_key, latest_date, float(metric_value)),
            4,
        )

    items: list[dict[str, object]] = []
    for country in ordered_countries:
        country_metrics = metrics_by_country.get(country, {})
        items.append(
            {
                "country": country,
                "code": SUPPLY_COUNTRY_CODE_MAP.get(country, _fallback_country_code(country)),
                "sourceLabel": source_by_country.get(country),
                "production": country_metrics.get("production"),
                "consumption": country_metrics.get("products_supplied"),
                "imports": country_metrics.get("imports"),
                "exports": country_metrics.get("exports"),
                "inventory": country_metrics.get("stock_change"),
                "refineryUtilization": country_metrics.get("refinery_utilization"),
            }
        )

    return {
        "dataset_code": SUPPLY_COUNTRY_DATASET_CODE,
        "as_of": latest_date,
        "items": items,
    }


def get_supply_ports_data(
    db_path: Path | str = DEFAULT_DB_PATH,
) -> dict[str, object]:
    init_market_database(db_path)
    with get_connection(db_path) as connection:
        rows = connection.execute(
            """
            SELECT *
            FROM port_reference_data
            WHERE dataset_code = ?
            ORDER BY country_name ASC, main_port_name ASC, id ASC
            """,
            (SUPPLY_PORT_DATASET_CODE,),
        ).fetchall()

    items = []
    for row in rows:
        items.append(
            {
                "id": str(row["wpi_number"] or row["id"]),
                "name": _translate_port_name(
                    row["main_port_name"],
                    row["alternate_port_name"],
                ),
                "nameOriginal": row["main_port_name"],
                "latitude": float(row["latitude"]),
                "longitude": float(row["longitude"]),
                "channelDepth": _format_port_depth_display(row["channel_depth_m"]),
                "anchorageDepth": _format_port_depth_display(row["anchorage_depth_m"]),
                "cargoPierDepth": _format_port_depth_display(row["cargo_pier_depth_m"]),
                "oilTerminalDepth": _format_port_depth_display(row["oil_terminal_depth_m"]),
                "harborSize": _translate_port_size(row["harbor_size"]),
                "harborType": _translate_port_type(row["harbor_type"]),
                "shelterAfforded": _translate_port_shelter(row["shelter_afforded"]),
                "waterBody": _translate_port_water_body(row["world_water_body"]),
                "detailLabel": _build_port_detail_label(row),
                "countryName": row["country_name"],
                "berthedVessels": [],
            }
        )

    return {
        "dataset_code": SUPPLY_PORT_DATASET_CODE,
        "total": len(items),
        "items": items,
    }


def count_news_rows(
    db_path: Path | str = DEFAULT_DB_PATH,
) -> int:
    init_market_database(db_path)
    with get_connection(db_path) as connection:
        row = connection.execute(
            """
            SELECT COUNT(*) AS count
            FROM market_news_items
            WHERE dataset_code = ?
            """,
            (NEWS_DATASET_CODE,),
        ).fetchone()
    return int(row["count"]) if row else 0


def bootstrap_news_database(
    source_path: Path | str | None = None,
    db_path: Path | str = DEFAULT_DB_PATH,
) -> dict[str, object] | None:
    workbook_path = Path(source_path) if source_path is not None else SAMPLE_DATA_DIR / "news_extraction.xlsx"
    if count_news_rows(db_path) > 0:
        return None
    if not workbook_path.exists():
        return None
    return import_news_extraction_workbook(workbook_path, db_path=db_path)


def get_news_items(
    section_key: str | None = None,
    limit: int | None = None,
    db_path: Path | str = DEFAULT_DB_PATH,
) -> list[dict[str, object]]:
    init_market_database(db_path)
    normalized_section_key = str(section_key or "").strip()
    params: list[object] = [NEWS_DATASET_CODE]
    where_clauses = ["dataset_code = ?"]
    if normalized_section_key and normalized_section_key.upper() != "HIGHLIGHT":
        where_clauses.append("section_key = ?")
        params.append(normalized_section_key)

    query = f"""
        SELECT id,
               title,
               source_label,
               published_at,
               theme_raw,
               section_key,
               section_name,
               market_impact_score
        FROM market_news_items
        WHERE {" AND ".join(where_clauses)}
        ORDER BY published_at DESC, COALESCE(market_impact_score, 0) DESC, id DESC
    """
    if limit is not None and limit > 0:
        query += " LIMIT ?"
        params.append(limit)

    with get_connection(db_path) as connection:
        rows = connection.execute(query, params).fetchall()

    return [
        {
            "id": str(row["id"]),
            "title": str(row["title"]),
            "source": str(row["source_label"]),
            "published_at": str(row["published_at"]),
            "theme": str(row["theme_raw"]),
            "section_key": str(row["section_key"]),
            "section_name": str(row["section_name"]),
            "impact": float(row["market_impact_score"]) if row["market_impact_score"] is not None else None,
        }
        for row in rows
    ]
